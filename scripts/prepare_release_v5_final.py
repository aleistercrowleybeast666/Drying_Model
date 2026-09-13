"""Prepare independent source-only V5 packages, verifying every seeded artifact."""
import argparse
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
NAME='A题_药材烘干模型'
OFFICIAL='01_四表正式数据_终端版'
GUI='02_完整功能_GUI版'
CORE=['materials','boundaries','operators','rk4','sampling','inputs','events','geometry','cases','stages','table_solver']
MODULES=['__init__','official_cli','official_progress','official_support',*CORE,'mesh','frozen_mesh',
    'table_reference','export','storage','diagnostics','judge_schedule','runtime']
TERMINAL_REQ='numpy>=2.0\nnumba>=0.61\nopenpyxl>=3.1\npsutil>=6\n'
GUI_REQ='numpy>=2.0\nscipy>=1.13\nnumba>=0.61\nmatplotlib>=3.9\nopenpyxl>=3.1\npillow>=10\nimageio>=2.35\nPySide6>=6.7\npypdf>=5\npsutil>=6\niapws==1.5.5\n'


def Prepare_Hash(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def Prepare_Read(path):return json.loads(Path(path).read_text(encoding='utf-8'))


def Prepare_Write(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def Prepare_Copy(source,target):
    source=Path(source);target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
    if target.exists():
        if Prepare_Hash(source)!=Prepare_Hash(target):raise ValueError('SEED_PATH_COLLISION: '+str(target))
    else:shutil.copy2(source,target)
    if source.samefile(target):raise ValueError('PACKAGE_PHYSICAL_LINK')


def Prepare_Entry(gui=False,backend=False):
    prefix='''import os
from pathlib import Path
import sys
sys.dont_write_bytecode=True
os.environ['PYTHONDONTWRITEBYTECODE']='1'
'''
    prefix+="root=Path(__file__).resolve()."+('parents[1]' if backend else 'parent')+'\n'
    prefix+="sys.path.insert(0,str(root/'dependencies/src'))\n"
    prefix+="os.environ['DRYING_MODEL_ROOT' if "+str(gui or backend)+" else 'DRYING_OFFICIAL_ROOT']=str(root)\n"
    prefix+="os.environ['NUMBA_CACHE_DIR']=str(root/"+repr('work/recompute/numba_cache' if gui or backend else 'work/numba_cache')+")\n"
    prefix+="os.environ['DRYING_JUDGE_FROZEN_MESH']='1'\n"
    prefix+="for name in ['OMP_NUM_THREADS','MKL_NUM_THREADS','OPENBLAS_NUM_THREADS','NUMEXPR_NUM_THREADS','NUMBA_NUM_THREADS']:os.environ[name]='1'\n"
    names=['numpy','numba','openpyxl','psutil']+(['scipy','matplotlib','PIL','imageio','PySide6','pypdf','iapws'] if gui or backend else [])
    prefix+='''import importlib.util
missing=[name for name in '''+repr(names)+''' if importlib.util.find_spec(name) is None]
if missing:
    print('缺少运行库：'+ '、'.join(missing)+'\\n请运行本目录中的“安装依赖.bat”\\n推荐 Python 3.11（64 位）')
    raise SystemExit(1)
'''
    if backend:return prefix+"from drying.judge_pipeline import Judge_Main\nraise SystemExit(Judge_Main(root))\n"
    if not gui:return prefix+"from drying.official_cli import Official_Main\nraise SystemExit(Official_Main(root))\n"
    return prefix+'''from PySide6.QtWidgets import QApplication
from drying.gui.judge_window import JudgeWindow
app=QApplication.instance() or QApplication(sys.argv)
window=JudgeWindow(root)
window.show()
if '--smoke' in sys.argv:
    import json
    from PySide6.QtCore import QTimer
    def Main_CheckReady():
        if window.catalog_panel.scan is not None:
            QTimer.singleShot(50,Main_CheckReady);return
        print(json.dumps(dict(gui_started=True,catalog_items=len(window.checks),snapshot=window.catalog_panel.snapshot),ensure_ascii=False),flush=True)
        window.close();app.quit()
    QTimer.singleShot(50,Main_CheckReady)
    QTimer.singleShot(60000,app.quit)
raise SystemExit(app.exec())
'''


def Prepare_Doc(text,gui=False,top=False):
    # Preserve user wording, replacing only obsolete entry/path/status statements.
    for old in ['药材烘干模型_四表正式数据.exe','药材烘干模型_GUI.exe','药材烘干模型_原题表格复算.exe']:
        text=text.replace(old,'main.py')
    text=text.replace('安装Python依赖.bat','安装依赖.bat').replace('dependencies/application/code','dependencies').replace('dependencies/application/data','dependencies/data')
    text=text.replace('双击','运行').replace('无需安装 Python','推荐 Python 3.11（64 位）')
    text=text.replace('同名 .py','main.py').replace('D 未运行不声称验证 PASS；','')
    text='\n'.join(line for line in text.splitlines() if not any(s in line for s in ['PyInstaller','本轮只进行少量','本轮按要求未执行','本轮按要求减少验证']))
    text=text.replace('“全选”只选择当前可用项，跳过禁用项','顶部“全选（A+B+C+D）”安排全部任务，先生成 A/B，再执行 C/D；模式内全选仍只选择当前可用项')
    text=text.replace('剩余时间不确定时暂保持并提示校准','剩余时间不确定时提示校准，结构进度继续更新')
    text=text.replace('D · 验证（建议不做）','D · 验证（按需选择）').replace('D · 验证（耗时最长，建议不做）','D · 验证（计算量较大，可按需选择）')
    if '推荐 Python 3.11（64 位）' not in text:text+='\n\n推荐 Python 3.11（64 位）。'
    text+='\n若提示缺少 Python 库，请运行对应目录中的“安装依赖.bat”。\n运行：python main.py\n'
    if gui:
        text=text.replace('无需预先保留结果或缓存；首次运行自动创建目录。','发布目录已附带当前核验有效的科学缓存和成果；可直接查看已有结果，也可选择重新计算/绘图。')
        text+='\nA = 四表正式数据；B = 创新实验数据；C = 绘图；D = 按需数值验证（计算量较大）。\n'
    elif not top:
        if '2.8' not in text:text+='\n用户测试机完整四表约 2.8 min；不同电脑和后台负载下耗时会变化，仅供参考。\n'
        if 'ETA' not in text:text+='运行时显示真实结构进度、已用时间和预计剩余时间（ETA），状态不更新时进度不会自行增长。\n'
        text+='首次运行仍会真实计算和核验四表；预置 Excel 不视为计算缓存命中。\n'
    text+='\n题目原始 PDF 与附件位于终端版、GUI 版各自的“原始题目数据”文件夹，两版均可单独复制使用。\n' if top else '\n题目原始 PDF 与附件位于本目录的“原始题目数据”文件夹。\n'
    return text+'\n'


def Prepare_RawInputs(package):
    files={}
    for path in sorted((ROOT/'data/raw').rglob('*')):
        if not path.is_file():continue
        relative=path.relative_to(ROOT/'data/raw')
        visible=Path(*relative.parts[1:]) if relative.parts[0]=='附件' else relative
        target=package/'原始题目数据'/visible
        Prepare_Copy(path,target)
        digest=Prepare_Hash(path)
        if Prepare_Hash(target)!=digest:raise ValueError('RAW_COPY_HASH_MISMATCH')
        files[relative.as_posix()]=dict(path=visible.as_posix(),sha256=digest)
    inputs=Prepare_Read(ROOT/'data/input_manifest.json')
    inputs.update(source='原始题目数据',source_base='package root',raw_files=files)
    for row in inputs['templates'].values():row['path']=Path(row['path']).as_posix()
    Prepare_Write(package/'dependencies/data/input_manifest.json',inputs)
    return files


def Prepare_Static(package,gui):
    package.mkdir(parents=True);dep=package/'dependencies'
    modules=list((ROOT/'src/drying').rglob('*.py')) if gui else [ROOT/'src/drying'/f'{name}.py' for name in MODULES]
    for path in modules:Prepare_Copy(path,dep/'src/drying'/path.relative_to(ROOT/'src/drying'))
    allowed=None if gui else {'default.toml','stage_schedule.json','stage_solver_reference.json','frozen_mesh','table_reference'}
    for path in (ROOT/'configs').rglob('*'):
        relative=path.relative_to(ROOT/'configs')
        if path.is_file() and path.suffix in ['.json','.toml','.npz','.npy'] and (allowed is None or relative.parts[0] in allowed):Prepare_Copy(path,dep/'configs'/relative)
    Prepare_RawInputs(package)
    Prepare_Copy(ROOT/'data/inputs.npz',dep/'data/inputs.npz')
    (package/'main.py').write_text(Prepare_Entry(gui),encoding='utf-8')
    if gui:(dep/'backend.py').write_text(Prepare_Entry(backend=True),encoding='utf-8')
    (package/'requirements.txt').write_text(GUI_REQ if gui else TERMINAL_REQ,encoding='utf-8')
    bat='''@echo off
chcp 65001 >nul
cd /d "%~dp0"
py -3.11 -c "import sys" >nul 2>&1
if not errorlevel 1 (
    py -3.11 -m pip install -r requirements.txt
) else (
    echo 推荐 Python 3.11（64 位），当前尝试使用 python。
    python -m pip install -r requirements.txt
)
set "install_result=%errorlevel%"
if not "%install_result%"=="0" echo 安装失败，错误码 %install_result%，请查看上方错误。
pause
exit /b %install_result%
'''
    (package/'安装依赖.bat').write_bytes(bat.replace('\n','\r\n').encode('utf-8'))
    original=ROOT/'release_v5'/NAME/(GUI if gui else OFFICIAL)
    for name in ['README.md','说明.txt']:
        (package/name).write_text(Prepare_Doc((original/name).read_text(encoding='utf-8'),gui),encoding='utf-8')
    for name in ['results','work','logs']:(package/name).mkdir()


def Seed_CheckTables(repo):
    import numpy as np
    from openpyxl import load_workbook
    from drying.export import Export_Readback
    reference=Prepare_Read(repo/'configs/table_reference/manifest.json');rows=[]
    for q in range(1,5):
        spec=reference['questions'][str(q)];data=repo/'configs/table_reference'/spec['file']
        if Prepare_Hash(data)!=spec['sha256']:raise ValueError('REFERENCE_HASH_MISMATCH')
        template=next((repo/'data/raw').rglob(f'result{q}.xlsx'))
        book=load_workbook(template,read_only=True);names=book.sheetnames;book.close()
        with np.load(data) as arrays:
            expected=[arrays['temperature_C'],arrays['moisture']] if q<=2 else [arrays['moisture']]
            columns=[round(i*.1,1) for i in range(21)]+(['药材表面'] if q==4 else [])
            Export_Readback(repo/f'results/tables/result{q}.xlsx',names,arrays['time_s'],expected,columns)
        rows.append(dict(question=q,status='PASS',sha256=Prepare_Hash(repo/f'results/tables/result{q}.xlsx')))
    return rows


def Seed_GuiScientificSnapshot(repo_root,final_gui_root):
    from drying.artifact_contract import Artifact_ReadManifest,Artifact_GetStatus,Artifact_GetIdentity,Artifact_Seal
    from drying.pipeline_plan import Catalog_Read
    repo=Path(repo_root).resolve();target=Path(final_gui_root).resolve();identity=Artifact_GetIdentity(repo);hashes={};pending={};skipped=[];copied={}
    for family in ['official','innovation','validation']:
        for key in Artifact_ReadManifest(repo,family+'.seed').get('entries',{}):
            check=Artifact_GetStatus(repo,key,True,hashes,identity)
            if check['complete'] and (family!='validation' or not key.startswith('validation.2d.') or check['entry'].get('evidence_complete')):pending[key]=check['entry']
            else:skipped.append(dict(key=key,reason=check.get('reason','raw validation evidence unavailable')))
    while pending:
        ready=[k for k,e in pending.items() if set(e.get('dependencies',[]))<=copied.keys()]
        if not ready:
            skipped += [dict(key=k,reason='dependency not migrated') for k in pending];break
        for key in ready:
            entry=pending.pop(key);runtime=(repo/entry.get('runtime','.')).resolve();paths=[]
            for name in entry['paths']:
                path=(repo/name).resolve()
                if not path.is_relative_to(repo):raise ValueError('SEED_PATH_OUTSIDE_REPO')
                if path.is_relative_to(repo/'results'):relative=path.relative_to(repo)
                else:
                    relative=path.relative_to(runtime)
                    if relative.parts[0]=='data':relative=Path('dependencies')/relative
                    if relative.parts[0] not in ['work','dependencies']:raise ValueError('SEED_NONSCIENTIFIC_PATH: '+name)
                if any(s in ['workers','tasks','receipts','numba_cache','archived','forced'] for s in relative.parts):raise ValueError('SEED_FORBIDDEN_PATH: '+name)
                if relative.as_posix()=='dependencies/data/input_manifest.json':
                    if Prepare_Read(path)['hash']!=Prepare_Read(target/relative)['hash']:raise ValueError('SEED_INPUT_IDENTITY_MISMATCH')
                else:Prepare_Copy(path,target/relative)
                paths.append(relative.as_posix())
            reserved={'key','producer','scientific_identity','dependencies','paths','sha256','complete','seal','dependency_seals'}
            metadata={k:v for k,v in entry.items() if k not in reserved};metadata['runtime']='.'
            copied[key]=Artifact_Seal(target,key,entry['producer'],paths,entry.get('dependencies',[]),metadata)
    publication=Prepare_Read(repo/'work/release/publication_manifest.json').get('publication_sha256',{})
    for name,digest in publication.items():
        relative=Path(name)
        allowed=name in ['results/paper_facts.json','results/paper_facts.md'] or (len(relative.parts)>2 and relative.parts[1] in ['q1','q2','q3','q4','animations']) or name.startswith(('results/studies/figures/','results/studies/animations/'))
        if not allowed or relative.suffix not in ['.png','.gif','.csv','.md','.json']:continue
        if relative.name.endswith('_summary.json'):continue
        path=repo/name
        if path.is_file() and Prepare_Hash(path)==digest:Prepare_Copy(path,target/name)
    for row in Catalog_Read(target).values():
        if row['mode']!='C':continue
        key=row['selection_key'];path=target/row['path']
        if path.is_file() and all(Artifact_GetStatus(target,k)['complete'] for k in row['requires']):
            copied[key]=Artifact_Seal(target,key,'seed_verified_existing_figure',[row['path']],row['requires'],dict(runtime='.',pde_solves=0,source_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=repo,text=True).strip()))
    counts={f:sum(k.startswith(f+'.') for k in copied) for f in ['official','innovation','plot','validation']}
    receipt=dict(copied=counts,skipped_stale_count=len(skipped),skipped=skipped,source_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=repo,text=True).strip(),source_scientific_identity=identity,target_scientific_identity=Artifact_GetIdentity(target))
    Prepare_Write(target/'results/data/seed_manifest.json',receipt)
    return receipt


def Prepare_Manifest(package,gui):
    files={p.relative_to(package).as_posix():Prepare_Hash(p) for p in package.rglob('*') if p.is_file() and p.relative_to(package).parts[0] not in ['work','results','logs'] and p.name!='package_manifest.json'}
    value=dict(package='full_gui' if gui else 'official_only',format='python_source',created_at=datetime.now().astimezone().isoformat(),files=files,
        official_core_sha256={name:files[f'dependencies/src/drying/{name}.py'] for name in CORE})
    Prepare_Write(package/'package_manifest.json',value);return value


def Prepare_CheckFinal(target):
    forbidden=['tests','scripts','__pycache__','.git','.github','.pytest_cache']
    for path in target.rglob('*'):
        rel=path.relative_to(target)
        if any(p in forbidden or p.startswith(('release_v','benchmark')) or 'pyinstaller' in p.lower() for p in rel.parts) or path.suffix in ['.exe','.spec','.pyc'] or path.name=='build.log':raise ValueError('FORBIDDEN_FINAL_FILE: '+str(rel))
    for name in [OFFICIAL,GUI]:
        if any((target/name/'logs').iterdir()):raise ValueError('FINAL_LOGS_NOT_EMPTY')
        if (target/name/'dependencies/data/raw').exists():raise ValueError('RAW_HIDDEN_IN_DEPENDENCIES')
        raw=Prepare_Read(target/name/'dependencies/data/input_manifest.json')['raw_files']
        for source,row in raw.items():
            if Prepare_Hash(ROOT/'data/raw'/source)!=row['sha256'] or Prepare_Hash(target/name/'原始题目数据'/row['path'])!=row['sha256']:raise ValueError('RAW_COPY_HASH_MISMATCH')
    if any((target/OFFICIAL/'work').iterdir()):raise ValueError('TERMINAL_WORK_NOT_EMPTY')
    manifests=[Prepare_Read(target/name/'package_manifest.json') for name in [OFFICIAL,GUI]]
    expected={name:Prepare_Hash(ROOT/f'src/drying/{name}.py') for name in CORE}
    if not all(m['official_core_sha256']==expected for m in manifests):raise ValueError('CORE_DIVERGED')
    return dict(clean=True,logs_empty=True,terminal_work_empty=True,numerical_core_sha256=expected)


def Prepare_Main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'release_v5_final'/NAME)
    parser.add_argument('--resume',action='store_true',help='Resume interrupted preparation without deleting any files')
    args=parser.parse_args()
    target=args.output.resolve()
    if not target.is_relative_to(ROOT/'release_v5_final') or target.exists() and not args.resume:raise ValueError('NEW_FINAL_DIRECTORY_REQUIRED')
    if not target.exists():
        target.mkdir(parents=True)
        for gui,name in [(False,OFFICIAL),(True,GUI)]:Prepare_Static(target/name,gui)
        for name in ['请评委先看.md','请评委先看.txt']:
            (target/name).write_text(Prepare_Doc((ROOT/'release_v5'/NAME/name).read_text(encoding='utf-8'),top=True),encoding='utf-8')
    print('Checking the four current workbooks against frozen full-precision references',flush=True)
    tables=Seed_CheckTables(ROOT)
    for name in [OFFICIAL,GUI]:
        for q in range(1,5):Prepare_Copy(ROOT/f'results/tables/result{q}.xlsx',target/name/f'results/tables/result{q}.xlsx')
    print('Deep-verifying and copying scientific artifacts; no PDE solves',flush=True)
    seed=Seed_GuiScientificSnapshot(ROOT,target/GUI)
    for gui,name in [(False,OFFICIAL),(True,GUI)]:Prepare_Manifest(target/name,gui)
    receipt=dict(seed=seed,workbook_readback=tables,checks=Prepare_CheckFinal(target))
    Prepare_Write(ROOT/'work/release_v5/final_source_audit/preparation.json',receipt)
    print(json.dumps(dict(output=str(target),seed=seed['copied'],skipped=seed['skipped_stale_count'],clean=True),ensure_ascii=False))


if __name__=='__main__':Prepare_Main()
