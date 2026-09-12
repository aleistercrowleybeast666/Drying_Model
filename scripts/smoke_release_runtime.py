"""Bounded frozen-worker capability checks, never a production trajectory solve."""
from pathlib import Path
import io
import inspect
import json
import sys
import time
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))


def Smoke_Run():
    import numpy as np
    from drying.cases import Case_LoadConfig,Case_LoadInputs,Case_GetSourceHash
    from drying.materials import Material_Evaluate
    from drying.studies.mass_kernel import MassBalance_GetFunctionTree
    from drying.rk4 import Rk4_Advance
    from drying.studies.scheduler import Studies_GetPlan
    config=Case_LoadConfig(ROOT); inputs=Case_LoadInputs(ROOT)
    assert config['physics']['C0']==2.55 and len(inputs)==3
    source=Path(inspect.getsourcefile(Material_Evaluate.py_func)).resolve()
    assert source.is_relative_to(ROOT/'src'), source
    assert len(MassBalance_GetFunctionTree(Rk4_Advance.py_func)[0].body)>0
    assert np.isfinite(Material_Evaluate(4,301.15,2.55)).all()
    from openpyxl import Workbook,load_workbook
    stream=io.BytesIO();wb=Workbook();wb.active['A1']=57.6215;wb.save(stream);stream.seek(0)
    assert load_workbook(stream,data_only=True).active['A1'].value==57.6215
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    figure,ax=plt.subplots();ax.plot([0,1],[0,1]);picture=io.BytesIO();figure.savefig(picture,format='png');plt.close(figure)
    from PIL import Image
    picture.seek(0);Image.open(picture).verify()
    print(json.dumps(dict(status='PASS',root=str(ROOT),numerical_source=str(source),core_hash=Case_GetSourceHash(ROOT),
        source_inspection_for_mass_audit=True,excel_roundtrip=True,agg_png=True,production_pde_solves=0),ensure_ascii=False),flush=True)
    if len(sys.argv)>1:
        Smoke_CheckRoute(Path(sys.argv[1]))
    return 0


def Smoke_CheckRoute(release):
    from PySide6.QtCore import QProcess
    from PySide6.QtWidgets import QApplication
    from drying.gui.judge_window import JudgeWindow
    from drying.recompute import TASKS
    assert getattr(sys,'frozen',False), 'This acceptance check must run inside the release EXE'
    app=QApplication.instance() or QApplication([])
    for keys,dry_run in [(['q1','q23','q4'],True),(list(TASKS),True),(['verify'],False)]:
        window=JudgeWindow(release);window.Selection_Set(keys);window.dry_run=dry_run
        window.Task_Start()
        deadline=time.monotonic()+90
        while window.process.state()!=QProcess.ProcessState.NotRunning and time.monotonic()<deadline:
            app.processEvents();time.sleep(.01)
        app.processEvents()
        assert window.process.state()==QProcess.ProcessState.NotRunning,'GUI route timed out'
        assert window.process.exitCode()==0,window.logs.toPlainText()
        assert Path(window.process.program())==release/'药材烘干模型_完整离线复算.exe'
        if dry_run: assert 'DRY_RUN' in window.logs.toPlainText()
        else: assert window.outer.value()==1 and window.outer.maximum()==1
        print(json.dumps(dict(status='PASS',gui_to_cli=True,tasks=keys,dry_run=dry_run,
            program=window.process.program(),production_pde_solves=0),ensure_ascii=False),flush=True)
        window.close()


if __name__=='__main__':raise SystemExit(Smoke_Run())
