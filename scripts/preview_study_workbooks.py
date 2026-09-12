"""Small faithful opening-view copies for bounded-memory workbook visual QA."""
import json
import sys
from copy import copy
from pathlib import Path
from openpyxl import Workbook,load_workbook


def Preview_Prepare(root,source_records=None,technical_only=False):
    root=Path(root)
    if source_records is None:source_records=json.loads((root/'work/studies/diagnostics/workbook_readback.json').read_text(encoding='utf-8'))['outputs']
    if technical_only:source_records=[r for r in source_records if '/tables/' in r['path']]
    output=root/'work/studies/diagnostics/workbook_previews';output.mkdir(parents=True,exist_ok=True);records=[]
    for number,record in enumerate(source_records):
        source=load_workbook(root/record['path']);preview=Workbook();preview.remove(preview.active);sheets=[]
        for original in source:
            if technical_only and original.title not in ['Geometry_Property_Cross','Kinetics','Kinetics_Summary']:continue
            row_limit=40 if technical_only and original.title=='Kinetics_Summary' else 12
            sheet=preview.create_sheet(original.title);sheet.sheet_view.showGridLines=original.sheet_view.showGridLines
            for row in original.iter_rows(max_row=min(original.max_row,row_limit)):
                for cell in row:
                    copied=sheet.cell(cell.row,cell.column,cell.value)
                    for name in ['font','fill','border','alignment','protection']:setattr(copied,name,copy(getattr(cell,name)))
                    copied.number_format=cell.number_format
            for name,dim in original.column_dimensions.items():sheet.column_dimensions[name].width=dim.width
            for row in range(1,min(original.max_row,row_limit)+1):sheet.row_dimensions[row].height=original.row_dimensions[row].height
            sheets.append(dict(name=sheet.title,columns=original.max_column,rows=min(original.max_row,row_limit),
                first_column=15 if technical_only and original.title=='Kinetics' else 1,column_block=4 if technical_only else 8,row_block=12 if technical_only else row_limit))
        path=output/f'book_{number}.xlsx';preview.save(path);preview.close();source.close()
        records.append(dict(source=record['path'],path=str(path),sheets=sheets))
    (output/'index.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--technical-only',action='store_true')
    Preview_Prepare(Path(__file__).resolve().parents[1],technical_only=parser.parse_args().technical_only)
