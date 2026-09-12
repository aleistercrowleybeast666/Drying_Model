"""Acceptance checks for final study artifacts and their immutable provenance."""
import argparse
import json
import sys
from pathlib import Path
import numpy as np
from PIL import Image
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from drying.studies.baseline import Baseline_ReadJson,Baseline_Check,Baseline_HashFile
from drying.studies.plot_contract import StudyPlot_ReadManifest,StudyPlot_Load,StudyPlot_Find
from drying.storage import Storage_WriteJson


def Studies_Check(root,allow_incomplete=False):
    root=Path(root);baseline=Baseline_Check(root);manifest=StudyPlot_ReadManifest(root)
    index=Baseline_ReadJson(root/'results/studies/study_index.json');problems=[]
    if index.get('current_complete_count')!=38:problems.append('not all 38 planned experiment trajectories complete')
    if len(manifest['checks'])!=21:problems.append('not all 21 independent spatial/time comparisons available')
    if len(manifest.get('thermal_audits',[]))!=27:problems.append('not all 27 thermal trajectory balance audits available')
    assert {entry['question'] for entry in manifest['end_effects']}=={'Q1','Q2','Q3','Q4'}
    for check in manifest['checks']:
        if not check['full_history_verified']:problems.append('reference lacks full initial history: '+check['reference_id'])
        if not check['reference_pairing_verified']:problems.append('reference scheme not paired to selected production: '+check['reference_id'])
        for event in check['events']:
            if event['source']!='genuine reintegration to exact common physical time':problems.append('event state not genuine')
        if check['mode']=='M00' and check['full_schedule_reference_passed'] is not True:problems.append('M00 full reference did not pass')
    for key,spec in manifest['specs'].items():
        status=manifest['statuses'][key];event=status.get('event')
        assert spec['fingerprint']==status['fingerprint'],key+' identity mismatch'
        if spec['kind']=='production' and spec['mode']!='M00':
            data=StudyPlot_Load(root,manifest['series'][key]);lo,hi=spec['water']['temperature_range_K']
            assert np.isfinite(data['profile']).all() and data['Tmin_K'].min()>=lo and data['Tmax_K'].max()<=hi
            assert np.min(data['profile'][:,1])>=0 and np.max(data['profile'][:,1])<=2.55+1e-10
        if event:
            assert event['left_max_C']>=spec['physics']['threshold']>event['right_max_C']
            assert event['report_max_C']<spec['physics']['threshold'] and event['width_s']<=.1+1e-12
        elif spec['case']!='q1':assert status.get('drying_time_h') is None
        if spec['kind']=='full_reference':
            first=spec['schedule'][0]
            assert first['t_start']==0 and status['simulated_time_s']==spec['schedule'][-1]['t_end']
            with np.load(sorted((root/'work/studies/experiments'/key).glob('chunk_*.npz'))[0]) as data:
                assert float(data['time_s'][0])==0
                assert np.all(data['fields'][0,0]==spec['physics']['T0_K']) and np.all(data['fields'][0,1]==spec['physics']['C0'])
    for entry in manifest['workbooks']:
        assert entry['readback']=='PASSED_ALL_CELLS' and Baseline_HashFile(root/entry['path'])==entry['sha256']
        if entry['path'].endswith('/supplement.xlsx'):
            mode=Path(entry['path']).parent.name
            required={'Config_Status','Q1_Temperature_C','Q1_Moisture_kg_kg','Q2_Temperature_C','Q2_Moisture_kg_kg',
                'Q3_Moisture_kg_kg','Q3_Temperature_C_extension','Q4_Moisture_kg_kg','Q4_Temperature_C_extension'}
            with_book=load_workbook(root/entry['path'],read_only=True,data_only=True)
            try:
                assert set(with_book.sheetnames)==required
                config=with_book['Config_Status'].values;headers=next(config);mode_col=headers.index('mode')
                assert all(row[mode_col]==mode for row in config)
            finally:with_book.close()
    assert len(manifest['workbooks'])==5
    assert len(list((root/'results/studies').rglob('*.xlsx')))==5
    if len(list((root/'results/studies/figures').glob('*.png')))!=9:problems.append('nine study figures not yet present')
    render_path=root/'work/studies/diagnostics/render_manifest.json'
    render=Baseline_ReadJson(render_path) if render_path.exists() else {}
    if render.get('manifest_seal')!=manifest['seal']:problems.append('figures need rendering against current numerical payload')
    gifs=[entry for entry in render.get('files',[]) if entry['path'].endswith('.gif')]
    if len(gifs)!=2:problems.append('both thermal GIFs not yet present')
    elif gifs[0]['frame_times_s']!=gifs[1]['frame_times_s']:problems.append('temperature/moisture GIF calendars differ')
    for entry in gifs:
        with Image.open(root/entry['path']) as gif:
            assert gif.n_frames==200 and gif.size==(1800,900)
            for i in range(gif.n_frames):gif.seek(i);gif.convert('RGB').load()
        assert entry['shared_physical_time'] and entry['frame_times_s'][-1]==259200.
        assert not entry['missing'],'missing thermal mode in final GIF'
        if entry['field']=='thermal_moisture':assert entry['fixed_color_limits']==[0,2.55]
    for mode in ['M00','M10','M01','M11']:
        try:
            entry=StudyPlot_Find(manifest,'q4',mode);data=StudyPlot_Load(root,entry)
        except FileNotFoundError:problems.append(mode+' Q4 production unavailable');continue
        assert data['radius_m'][-1]<data['radius_m'][0]
        state=manifest['statuses'][entry['experiment_id']]
        if state.get('event'):assert data['Cmax'][-1]<state['event']['report_max_C']
    numerical_failures=[c['reference_id'] for c in manifest['checks'] if c['status']!='PASS']
    numerical_failures += [a['experiment_id'] for a in manifest.get('thermal_audits',[]) if a['status']!='PASS']
    numerical_failures += [a['case']+'_M00_remesh' for a in manifest.get('remesh_front_audits',[]) if a['status']!='PASS']
    technical=None
    if manifest.get('technical_extension'):
        from drying.studies.technical_audit import Technical_CheckFinal
        technical=Technical_CheckFinal(root,manifest)
        numerical_failures+=technical['problems']
    gui_path=root/'work/studies/diagnostics/gui_cli_check.json'
    gui=Baseline_ReadJson(gui_path) if gui_path.exists() else dict(status='NOT_RUN',manual_interaction=False)
    if gui.get('status')=='PASS' and gui.get('app_sha256')!=Baseline_HashFile(root/'app.py'):gui=dict(gui,status='STALE_TEST')
    result=dict(status='INCOMPLETE' if problems else 'NUMERICAL_CHECK_FAILED' if numerical_failures else 'PASS',baseline=baseline,artifact_issues=problems,
        current_trajectories=index.get('current_complete_count'),expected_trajectories=38,checks=len(manifest['checks']),
        technical_extension=technical,numerical_failures=numerical_failures,thermal_audits=manifest.get('thermal_audits',[]),
        gui=gui,gui_scope='hidden Tk configuration and actual read-only child process; no full manual clicking session claimed',
        notes=['2D remains PARTIAL','fixed-radius appendix-4 control may be NOT_DRY_WITHIN_72H','legacy fixed grid is diagnostic only'])
    Storage_WriteJson(root/'work/studies/validation/acceptance.json',result)
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if (problems or numerical_failures) and not allow_incomplete:raise SystemExit(1)
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--allow-incomplete',action='store_true')
    Studies_Check(ROOT,parser.parse_args().allow_incomplete)
