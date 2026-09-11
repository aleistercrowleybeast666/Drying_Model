import json
import shutil
import stat
import zipfile
from pathlib import Path, PurePosixPath
import numpy as np
from numba import njit
from openpyxl import load_workbook
from .storage import Storage_HashFiles, Storage_WriteJson, Storage_WriteArray


def Input_ExtractZip(source, destination):
    destination = Path(destination).resolve()
    with zipfile.ZipFile(source) as archive:
        for entry in archive.infolist():
            name = entry.filename
            if not entry.flag_bits & 0x800:
                try:
                    name = name.encode('cp437').decode('gbk')
                except (UnicodeError, LookupError):
                    pass
            relative = PurePosixPath(name.replace('\\', '/'))
            if relative.is_absolute() or '..' in relative.parts or any(':' in p for p in relative.parts):
                raise ValueError('INPUT_SCHEMA_ERROR: unsafe ZIP path')
            if stat.S_ISLNK(entry.external_attr >> 16):
                raise ValueError('INPUT_SCHEMA_ERROR: ZIP symlink rejected')
            parts = relative.parts
            positions = [i for i, part in enumerate(parts) if part == 'A题']
            if not positions or entry.is_dir() or relative.name.startswith('~$'):
                continue
            target = destination.joinpath(*parts[positions[0]+1:]).resolve()
            if not target.is_relative_to(destination):
                raise ValueError('INPUT_SCHEMA_ERROR: ZIP traversal')
            if target.exists():
                if target.read_bytes() != archive.read(entry):
                    raise ValueError('CACHE_MISMATCH: existing raw data differs')
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(entry) as src, target.open('xb') as dst:
                shutil.copyfileobj(src, dst)


def Input_ReadTable(path, headers, count):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.values)
    workbook.close()
    observed = [str(x).strip() for x in rows[0]]
    if set(headers) != set(observed):
        raise ValueError(f'INPUT_SCHEMA_ERROR: {path}: {observed}')
    columns = [observed.index(h) for h in headers]
    rows = [r for r in rows[1:] if any(v is not None for v in r)]
    result = np.asarray([[r[j] for j in columns] for r in rows], dtype=np.float64)
    if result.shape != (count, len(headers)) or not np.isfinite(result).all():
        raise ValueError(f'INPUT_VALUE_INVALID: {path}: shape/finite values')
    return result


def Input_Prepare(root, source):
    root = Path(root)
    raw = root/'data/raw'
    raw.mkdir(parents=True, exist_ok=True)
    source = Path(source)
    if source.suffix.lower() == '.zip':
        Input_ExtractZip(source, raw)
    elif source.is_dir():
        for path in source.rglob('*'):
            if path.is_file() and not path.name.startswith('~$') and path.suffix.lower() in ('.pdf', '.xlsx'):
                target = raw/path.relative_to(source)
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.exists() and target.read_bytes() != path.read_bytes():
                    raise ValueError('CACHE_MISMATCH: refusing to replace raw input')
                if not target.exists():
                    shutil.copy2(path, target)
    else:
        raise FileNotFoundError(f'INPUT_MISSING: {source}')
    env_path = next(raw.rglob('附件1.xlsx'), None)
    radius_path = next(raw.rglob('附件2.xlsx'), None)
    pdf_path = next(raw.rglob('A题.pdf'), None)
    if not all((env_path, radius_path, pdf_path)):
        raise FileNotFoundError('INPUT_MISSING: A题.pdf / 附件1.xlsx / 附件2.xlsx')
    env = Input_ReadTable(env_path, ['时间', '温度', '水分浓度'], 241)
    radius = Input_ReadTable(radius_path, ['时间', '半径'], 145)
    if not np.array_equal(env[:, 0], np.arange(241)*60) or not np.array_equal(radius[:, 0], np.arange(145)*1800):
        raise ValueError('INPUT_VALUE_INVALID: time nodes')
    if not np.allclose(radius[[0, -1], 1], [2, 1.198], atol=1e-12, rtol=0):
        raise ValueError('UNIT_CHECK_FAILED: radius endpoints in cm')
    if np.any(radius[:, 1] <= 0) or np.any(env[:, 2] <= 0):
        raise ValueError('INPUT_VALUE_INVALID: nonpositive inputs')
    tail = env[(env[:, 0] >= 10800) & (env[:, 0] <= 14400), 1:].mean(axis=0)
    if not np.allclose(tail, [49.99893442622951, 0.04998754098360656], rtol=0, atol=1e-11):
        raise ValueError(f'SOURCE_FORMULA_MISMATCH: tail regression {tail}')
    templates = {}
    for q in range(1, 5):
        path = next(raw.rglob(f'result{q}.xlsx'), None)
        if path is None:
            raise FileNotFoundError(f'INPUT_MISSING: result{q}.xlsx template')
        wb = load_workbook(path, read_only=True)
        expected = ['温度', '水分浓度'] if q <= 2 else ['Sheet1']
        if wb.sheetnames != expected:
            raise ValueError(f'INPUT_SCHEMA_ERROR: {path} sheet names')
        templates[str(q)] = dict(path=str(path.relative_to(root)),
            sheets={s.title: dict(header=list(next(s.values)), rows=s.max_row, columns=s.max_column) for s in wb})
        wb.close()
    from pypdf import PdfReader
    text = '\n'.join(page.extract_text() for page in PdfReader(pdf_path).pages)
    (root/'data/problem_text.txt').write_text(text, encoding='utf-8')
    # PDF formula extraction is not a symbolic parser; its text is saved for audit.
    manifest = dict(source=str(source.resolve()), hash=Storage_HashFiles(raw.rglob('*.*')),
        environment_samples=241, tail_samples=61, T_tail_C=float(tail[0]), H_tail=float(tail[1]),
        radius_samples=145, R0_m=0.02, R_end_m=0.01198, templates=templates,
        source_formula_review='PDF appendices 2/3/4 visually and textually checked against prompt')
    env[:, 1] += 273.15
    radius[:, 1] *= 0.01
    tail[0] += 273.15
    Storage_WriteArray(root/'data/inputs.npz', environment=env, radius=radius, tail=tail)
    Storage_WriteJson(root/'data/input_manifest.json', manifest)
    return manifest


@njit(cache=True)
def Input_AtTime(t, environment, radius, tail, shrink, right_at_tail=False):
    if t < 0 or (shrink and t > radius[-1, 0] + 1e-8):
        raise ValueError('RADIUS_OUT_OF_RANGE / ENVIRONMENT_RANGE_ERROR')
    if t > 14400 or (right_at_tail and t >= 14400):
        Te, He = tail[0], tail[1]
    else:
        Te = np.interp(t, environment[:, 0], environment[:, 1])
        He = np.interp(t, environment[:, 0], environment[:, 2])
    R = np.interp(t, radius[:, 0], radius[:, 1]) if shrink else 0.02
    return Te, He, R
