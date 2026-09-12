"""Check that rendered opening views still match the final workbook cells/styles."""
from pathlib import Path
from copy import copy
import gc
import json
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[3]
folder = root / 'work/studies/diagnostics/workbook_previews'
records = json.loads((folder / 'index.json').read_text(encoding='utf-8'))
checked = 0
for record in records:
    source = load_workbook(root / record['source'])
    preview = load_workbook(record['path'])
    for name in source.sheetnames:
        original, rendered = source[name], preview[name]
        for row in original.iter_rows(max_row=min(original.max_row, 12)):
            for cell in row:
                other = rendered.cell(cell.row, cell.column)
                assert cell.value == other.value, (record['source'], name, cell.coordinate, 'value')
                assert cell.number_format == other.number_format
                for attribute in ['font', 'fill', 'border', 'alignment', 'protection']:
                    assert copy(getattr(cell, attribute)) == copy(getattr(other, attribute)), (name, cell.coordinate, attribute)
                checked += 1
        for column, dimension in original.column_dimensions.items():
            assert dimension.width == rendered.column_dimensions[column].width
        for row in range(1, min(original.max_row, 12) + 1):
            assert original.row_dimensions[row].height == rendered.row_dimensions[row].height
    source.close()
    preview.close()
    del source, preview
    gc.collect()
result = dict(status='PASS', workbooks=len(records), checked_preview_cells=checked,
              scope='all rendered first-12-row views equal final workbook values, units and visible cell styles')
(folder / 'final_source_check.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
print(json.dumps(result))
