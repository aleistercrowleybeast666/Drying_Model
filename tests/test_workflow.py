import json
import zipfile
from pathlib import Path
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from drying.cases import Case_LoadInputs
from drying.inputs import Input_ExtractZip
from drying.events import Event_Locate
from drying.sampling import Sampling_GetRadial
from drying.export import Export_WriteWorkbook, Export_Readback
from drying.storage import Storage_WriteArray, Storage_WriteJson


def test_real_attachment_tail_and_units():
    root=Path(__file__).resolve().parents[1]
    env,radius,tail=Case_LoadInputs(root)
    assert env.shape==(241,3) and radius.shape==(145,2)
    np.testing.assert_allclose(tail,[323.1489344262295,.04998754098360656],rtol=0,atol=1e-11)
    np.testing.assert_allclose(radius[[0,-1],1],[.02,.01198],rtol=0,atol=1e-14)


def test_spatial_export_outside_domain_and_true_surface():
    env=np.array([[0.,301.15,.02],[14400.,323.15,.05]])
    radius=np.array([[0.,.02],[259200.,.01198]])
    inputs=(env,radius,np.array([323.15,.05]))
    state=np.empty((2,40,1)); state[0]=320.; state[1]=.3
    values,valid,surface,R=Sampling_GetRadial(state,259200.,4,inputs,[0,.005,.01,.015,.02])
    assert valid.tolist()==[True,True,True,False,False]
    assert np.isnan(values[:,3:]).all() and R==.01198
    assert .05<surface[1]<.3


def test_event_right_endpoint_strictly_dry_and_report_reintegrated():
    env=np.array([[0.,301.15,.02],[14400.,323.15,.05]])
    radius=np.array([[0.,.02],[259200.,.01198]])
    inputs=(env,radius,np.array([323.15,.05]))
    state=np.empty((2,4,1)); state[0]=301.15; state[1]=.1501
    called=[]
    def Advance(old,left,right):
        called.append(right)
        result=old.copy(); result[1]=.151-.00001*right
        return (result,)
    event,report=Event_Locate(90.,110.,state,Advance,3,inputs)
    assert event['left_s']<=100<event['right_s']
    assert event['width_s']<=.1 and event['report_s']>=event['right_s']
    assert event['report_s'] in called and report[1].max()<.15
    assert event['left_max_C']>=.15 and event['right_max_C']<.15


def test_zip_traversal_and_original_preservation(tmp_path):
    malicious=tmp_path/'bad.zip'
    with zipfile.ZipFile(malicious,'w') as archive:
        archive.writestr('A题/../outside.txt','malicious')
    with pytest.raises(ValueError,match='unsafe'):
        Input_ExtractZip(malicious,tmp_path/'raw')
    good=tmp_path/'good.zip'
    with zipfile.ZipFile(good,'w') as archive:
        archive.writestr('A题/test.txt','original')
        archive.writestr('B题/test.txt','unrelated')
    Input_ExtractZip(good,tmp_path/'raw')
    assert (tmp_path/'raw/test.txt').read_text()=='original'
    (tmp_path/'raw/test.txt').write_text('user changed')
    with pytest.raises(ValueError,match='CACHE_MISMATCH'):
        Input_ExtractZip(good,tmp_path/'raw')


def test_excel_roundtrip_numeric_blank_and_four_decimals(tmp_path):
    template=tmp_path/'template.xlsx'
    wb=Workbook(); wb.active.title='Sheet1'; wb.active.append(['时间',0,'…','药材表面']); wb.save(template); wb.close()
    times=np.array([60.,120.,120.24]); values=np.array([[.2,np.nan,.1],[.16,np.nan,.07],[.1499999,np.nan,.06]])
    path=tmp_path/'result.xlsx'
    Export_WriteWorkbook(template,path,times,[values],[0.,2.,'药材表面'])
    Export_Readback(path,['Sheet1'],times,[values],[0.,2.,'药材表面'])
    assert not (tmp_path/'result.tmp.xlsx').exists()


def test_atomic_full_precision_storage(tmp_path):
    values=np.array([.1499999999991,np.nan])
    Storage_WriteArray(tmp_path/'cache.npz',values=values)
    with np.load(tmp_path/'cache.npz') as cache: np.testing.assert_array_equal(cache['values'],values)
    Storage_WriteJson(tmp_path/'status.json',dict(drying_time_h=None,status='NOT_DRY_WITHIN_72H'))
    assert json.loads((tmp_path/'status.json').read_text())['drying_time_h'] is None


def test_vscode_workspace_valid():
    root=Path(__file__).resolve().parents[1]
    workspace=json.loads((root/'A_Drying.code-workspace').read_text(encoding='utf-8'))
    assert workspace['folders'][0]['path']=='.'
    assert '.venv/Scripts/python.exe' in workspace['settings']['python.defaultInterpreterPath']
    assert len(workspace['tasks']['tasks'])>=6
