"""Portable, sealed per-experiment data contracts. Never starts a solver."""
from enum import Enum
import hashlib
import json
from pathlib import Path
from .storage import Storage_WriteJson


class ArtifactCheckResult(str,Enum):
    COMPLETE='COMPLETE'
    MISSING='MISSING'
    STALE='STALE'


def Artifact_GetIdentity(root):
    from .judge_progress import Progress_GetIdentity
    from .runtime import Runtime_GetCode
    identity=Progress_GetIdentity(root)
    code=Runtime_GetCode(Path(root))
    identity['experiment_sources']={name:Artifact_HashFile(code/'src/drying/studies'/name) for name in ['trajectory.py','thermal.py','water.py','cross_spec.py']}
    return identity


def Artifact_HashFile(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def Artifact_GetSeal(value):
    return hashlib.sha256(json.dumps({k:v for k,v in value.items() if k!='seal'},sort_keys=True,allow_nan=False).encode()).hexdigest()


def Artifact_GetManifestPath(root,key):
    family=key.split('.')[0]
    names={'official':'official_manifest','innovation':'innovation_manifest','validation':'validation_manifest','plot':'plot_manifest'}
    return Path(root)/'results/data'/(names[family]+'.json')


def Artifact_ReadManifest(root,key):
    path=Artifact_GetManifestPath(root,key)
    if not path.exists():return {}
    try:
        value=json.loads(path.read_text(encoding='utf-8'))
        return value if value.get('seal')==Artifact_GetSeal(value) else {}
    except (ValueError,OSError):return {}


def Artifact_GetStatus(root,artifact_key,deep=False,hashes=None,identity=None,manifests=None):
    root=Path(root).resolve()
    family=artifact_key.split('.')[0]
    manifest=Artifact_ReadManifest(root,artifact_key) if manifests is None else manifests.get(family,{})
    entry=manifest.get('entries',{}).get(artifact_key)
    if not entry:return dict(key=artifact_key,status=ArtifactCheckResult.MISSING,complete=False,reason='未注册数据清单')
    try:
        if entry.get('scientific_identity')!=(Artifact_GetIdentity(root) if identity is None else identity):raise ValueError('科学身份已变化')
        if not entry.get('complete') or entry.get('seal')!=Artifact_GetSeal(entry):raise ValueError('清单未完成或封印失配')
        if not entry.get('sha256'):raise ValueError('数据清单为空')
        for key,seal in entry.get('dependency_seals',{}).items():
            source=Artifact_ReadManifest(root,key).get('entries',{}).get(key,{})
            if source.get('seal')!=seal:raise ValueError('前置数据已更新 '+key)
        hashes={} if hashes is None else hashes
        for name,digest in entry['sha256'].items():
            path=(root/name).resolve()
            if not path.is_relative_to(root) or not path.is_file():raise ValueError('缺少文件 '+name)
            if deep:
                if name not in hashes:hashes[name]=Artifact_HashFile(path)
                if hashes[name]!=digest:raise ValueError('文件哈希失配 '+name)
        return dict(key=artifact_key,status=ArtifactCheckResult.COMPLETE,complete=True,entry=entry)
    except (ValueError,OSError,KeyError) as error:
        return dict(key=artifact_key,status=ArtifactCheckResult.STALE,complete=False,reason=str(error))


def Artifact_Verify(root,artifact_key):return Artifact_GetStatus(root,artifact_key,deep=True)


def Artifact_Require(root,artifact_keys,error_code='PREREQUISITE_MISSING'):
    hashes={};entries={};missing=[]
    for key in dict.fromkeys(artifact_keys):
        status=Artifact_GetStatus(root,key,True,hashes)
        if status['complete']:entries[key]=status['entry']
        else:missing.append(dict(key=key,reason=status.get('reason',status['status'])))
    if missing:raise RuntimeError(error_code+': '+json.dumps(missing,ensure_ascii=False)+'；请先运行对应 A/B 数据任务，不会自动补算。')
    return entries


def Artifact_Seal(root,artifact_key,producer,paths,dependencies=(),metadata=None):
    root=Path(root).resolve();files={}
    for name in paths:
        path=(root/name).resolve()
        if not path.is_relative_to(root) or not path.is_file():raise ValueError('ARTIFACT_PATH_INVALID: '+str(name))
        files[path.relative_to(root).as_posix()]=Artifact_HashFile(path)
    if not files:raise ValueError('ARTIFACT_FILES_REQUIRED: '+artifact_key)
    identity=Artifact_GetIdentity(root)
    entry=dict(key=artifact_key,producer=producer,scientific_identity=identity,dependencies=list(dependencies),
        paths=list(files),sha256=files,complete=True,**(metadata or {}))
    entry['dependency_seals']={k:Artifact_ReadManifest(root,k).get('entries',{}).get(k,{}).get('seal') for k in dependencies}
    entry['seal']=Artifact_GetSeal(entry)
    manifest=Artifact_ReadManifest(root,artifact_key)
    if manifest.get('scientific_identity')!=identity:manifest={}
    entries=manifest.get('entries',{});entries[artifact_key]=entry
    family=artifact_key.split('.')[0]
    required=['official.'+c for c in ['q1','q23','q4']] if family=='official' else list(entries)
    value=dict(schema_version=1,scientific_identity=identity,input_hash=identity['input_hash'],
        numerical_core_hash=identity['numerical_source_hash'],frozen_mesh_hash=identity['frozen_mesh_hash'],
        schedule_hash=identity['frozen_schedule_hash'],entries=entries,required_entries=required,
        complete=all(entries.get(k,{}).get('complete') for k in required))
    if family=='official':
        value.update({c:entries.get('official.'+c,{}) for c in ['q1','q23','q4']})
        value['tables']={Path(n).name:h for e in entries.values() for n,h in e['sha256'].items() if n.endswith('.xlsx')}
    value['seal']=Artifact_GetSeal(value);Storage_WriteJson(Artifact_GetManifestPath(root,artifact_key),value)
    if family in ['official','innovation']:
        Storage_WriteJson(root/'work/datasets'/family/(artifact_key+'.json'),entry)
    return entry


def Artifact_GetCachePaths(runtime,status):
    runtime=Path(runtime);paths=[]
    if 'case_id' in status and (runtime/'work/cache'/status['case_id']).is_dir():
        for item in [status,*status.get('stages',[])]:
            folder=runtime/'work/cache'/item['case_id']
            paths += [p for p in folder.iterdir() if p.is_file() and (p.suffix=='.npz' or p.name=='status.json')]
    else:
        folder=runtime/'work/studies/experiments'/status['experiment_id']
        paths += [p for p in folder.iterdir() if p.is_file() and p.suffix in ['.npz','.json']]
    return paths


def Artifact_Hydrate(root,workspace,entries):
    """Link sealed numerical inputs into a task-private layout, never mutate them."""
    from .judge_resources import Resource_CopyFile
    root=Path(root);workspace=Path(workspace)
    for entry in entries.values():
        source_root=root/entry.get('runtime','work/recompute/runtime')
        for name in entry['paths']:
            source=root/name
            try:relative=source.relative_to(source_root)
            except ValueError:continue
            if relative.parts[0] in ['work','data']:Resource_CopyFile(source,workspace/relative)


def Artifact_ReadbackOfficial(root,entry):
    import numpy as np
    from openpyxl import load_workbook
    from .export import Export_Readback
    root=Path(root);runtime=root/entry['runtime'];inputs=json.loads((runtime/'data/input_manifest.json').read_text(encoding='utf-8'))
    for q in dict(q1=[1],q23=[2,3],q4=[4])[entry['case']]:
        with np.load(runtime/f'work/cache/exports/result{q}_full_precision.npz') as data:
            arrays=[data['temperature_C'],data['moisture']] if q<=2 else [data['moisture']]
            columns=[round(x*.1,1) for x in range(21)]+(['药材表面'] if q==4 else [])
            book=load_workbook(runtime/inputs['templates'][str(q)]['path'],read_only=True);names=book.sheetnames;book.close()
            Export_Readback(root/f'results/tables/result{q}.xlsx',names,data['time_s'],arrays,columns)
    return dict(status='PASS',pde_solves=0)
