"""Five supplementary workbooks, atomic writes and full numeric readback.

The four official templates and Export_WriteWorkbook/Export_Readback pipeline
are reused in work/studies before their verified sheets are assembled together.
"""
import json
import math
import os
from copy import copy
from pathlib import Path
import numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import CellIsRule
from ..export import Export_WriteWorkbook
from ..storage import Storage_WriteJson
from .baseline import Baseline_ReadJson,Baseline_HashFile


def Export_GetValue(value):
    if isinstance(value,np.generic):value=value.item()
    if isinstance(value,float) and not np.isfinite(value):return None
    if isinstance(value,(list,dict,tuple)):return json.dumps(value,ensure_ascii=False)
    return value


def Export_CopyCellStyle(source,target):
    # Style IDs belong to one workbook. Register the actual components when
    # copying a verified template sheet into a different workbook.
    for name in ['font','fill','border','alignment','protection']:
        setattr(target,name,copy(getattr(source,name)))
    target.number_format=source.number_format


def Export_GetColumnLabel(key):
    if key.endswith(('Cmax','Cmean','_max_C')) or key in ['max_abs_C','minimum_moisture','max_common_Cmax_difference']:
        return key+' / kg/kg'
    if key=='max_abs_T' or key.endswith('_max_T'):return key+' / K'
    return key


def Export_AddTable(wb,name,records):
    sheet=wb.create_sheet(name);sheet.sheet_view.showGridLines=False
    if not records:records=[dict(status='NOT_APPLICABLE_OR_INCOMPLETE',reason='See study_index.json for pending or failed prerequisite')]
    columns=list(dict.fromkeys(k for row in records for k in row));sheet.append([Export_GetColumnLabel(k) for k in columns])
    for row in records:sheet.append([Export_GetValue(row.get(k)) for k in columns])
    sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='28455D')
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    sheet.row_dimensions[1].height=32
    for col in sheet.columns:
        cells=list(col);key=str(cells[0].value)
        width=max(13,min(52,max(len(str(c.value or '')) for c in cells[:80])+2))
        if key in ['experiment_id','reference','control_id','source','reason','meaning','interpretation']:width=48
        if key=='source':width=64
        sheet.column_dimensions[cells[0].column_letter].width=width
        for cell in cells[1:]:
            cell.font=Font(name='Arial',size=10,color='202A35')
            cell.alignment=Alignment(vertical='center',wrap_text=isinstance(cell.value,str) and len(cell.value)>width)
            if isinstance(cell.value,float):cell.number_format='0.000000' if abs(cell.value)>=1e-4 or cell.value==0 else '0.0000E+00'
            if isinstance(cell.value,str) and len(cell.value)>width:
                sheet.row_dimensions[cell.row].height=max(sheet.row_dimensions[cell.row].height or 15,min(180,14*math.ceil(len(cell.value)/(width*.9))))
    return sheet


def Export_SaveVerified(root,wb,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);temp=path.with_suffix('.tmp.xlsx')
    # Expected values are taken before serialization; every saved cell is read back.
    expected={s.title:tuple(tuple(Export_GetValue(v) for v in row) for row in s.values) for s in wb}
    wb.save(temp);wb.close();loaded=load_workbook(temp,read_only=True,data_only=True);cells=0
    try:
        assert loaded.sheetnames==list(expected)
        for sheet in loaded:
            rows=expected[sheet.title];actual=tuple(sheet.values)
            assert len(actual)==len(rows)
            for a,b in zip(actual,rows):
                assert len(a)==len(b)
                for value,reference in zip(a,b):
                    if isinstance(reference,(float,int)) and not isinstance(reference,bool):
                        assert isinstance(value,(float,int)) and abs(value-reference)<=max(2e-12,abs(reference)*2e-15)
                    else:assert value==reference or (value is None and reference=='')
                    cells+=1
    finally:loaded.close()
    os.replace(temp,path)
    return dict(path=path.relative_to(root).as_posix(),sha256=Baseline_HashFile(path),sheets=list(expected),
        readback='PASSED_ALL_CELLS',checked_cells=cells)


def Export_GetTimeseries(root,manifest):
    from .analysis import Analysis_LoadSeries
    tables={k:[] for k in ['EndEffects','Fronts','Kinetics','DiffusionClock','Counterfactuals','ThermalSeries']}
    for key,entry in manifest['series'].items():
        if entry['kind']!='production':continue
        data=Analysis_LoadSeries(root,entry)
        # Exact cached sample rows: 60 s through 4h, 10 min thereafter, plus extrema/switches.
        t=data['time_s'];keep=((t<=14400)&(np.mod(t,60)<1e-8)) | ((t>14400)&(np.mod(t,600)<1e-8))
        keep[0]=True;keep[-1]=True
        for variable in ['loss_Cmean_s','Dmean_m2_s','Tmin_K']:keep[np.nanargmax(data[variable])]=True;keep[np.nanargmin(data[variable])]=True
        switches=np.flatnonzero(np.diff(data['stage_nr'])!=0)+1;keep[switches]=True
        keep|=np.isin(t,entry.get('exact_event_times_s',[]))
        for i in np.flatnonzero(keep):
            row=dict(case=entry['case'],mode=entry['mode'],time_s=float(t[i]))
            tables['Kinetics'].append(dict(row,Cmax=data['Cmax'][i],Cmean=data['Cmean'][i],loss_Cmax_kg_kg_s=data['loss_Cmax_s'][i],
                loss_Cmean_kg_kg_s=data['loss_Cmean_s'][i],D_volume_m2_s=data['Dmean_m2_s'][i],D_center_m2_s=data['Dcenter_m2_s'][i],
                radius_m=data['radius_m'][i],radius_rate_m_s=data['radius_rate_m_s'][i],max_tie_count=data['ties'][i,0],
                max_r_min_m=data['ties'][i,1],max_r_max_m=data['ties'][i,2]))
            tables['DiffusionClock'].append(dict(row,theta_V=data['theta_V'][i],theta_c=data['theta_c'][i],Cmax=data['Cmax'][i],
                **({f'{name}_{place}':data['drivers_'+place][i,j] for place in ['center','mid'] for j,name in enumerate(['BT','BC','BR'])} if 'drivers_center' in data else {}),
                **({f'identity_residual_{place}':data['drivers_residual_'+place][i] for place in ['center','mid']} if 'drivers_center' in data else {})))
            if entry['mode']=='M00':
                for k,c in enumerate([2.,1.,.5,.15]):
                    tables['Fronts'].append(dict(row,threshold_kg_kg=c,rf_m=data['fronts'][i,k,0],sf=data['fronts'][i,k,1],
                        wet_volume_fraction=data['fronts'][i,k,2],status=data['front_status'][i,k]))
    for entry in manifest['end_effects']:
        with np.load(root/entry['path']) as data:
            valid=np.flatnonzero(data['time_s']<=entry['time_range_s'][-1]);keep=valid[::10]
            important=entry.get('exact_event_times_s',[])+[v for k,v in entry.items() if k.startswith('peak_') and k.endswith('_time_s')]
            keep=np.unique(np.r_[keep,valid[-1],np.flatnonzero(np.isin(data['time_s'],important))])
            for i in keep:tables['EndEffects'].append(dict(case=entry['case'],question=entry['question'],**{k:Export_GetValue(data[k][i]) for k in data.files}))
    from .synthesis import Synthesis_Find
    for case in ['q23','q4']:
        keys=[Synthesis_Find(manifest,case,mode) for mode in ['M00','M10','M01','M11']]
        keys=[key for key in keys if key is not None]
        values={key:Analysis_LoadSeries(root,manifest['series'][key]) for key in keys}
        common=values[keys[0]]['time_s']
        for key in keys[1:]:common=np.intersect1d(common,values[key]['time_s'])
        selected=((common<=14400)&(np.mod(common,60)<1e-8))|((common>14400)&(np.mod(common,600)<1e-8))
        selected[0]=True;selected[-1]=True
        for key in keys:
            d=values[key];important=list(manifest['series'][key].get('exact_event_times_s',[]))
            for field in ['Tmin_K','loss_Cmean_s','loss_Cmax_s']:important += [d['time_s'][np.nanargmin(d[field])],d['time_s'][np.nanargmax(d[field])]]
            selected|=np.isin(common,important)
        for time in common[selected]:
            for key in keys:
                d=values[key];i=np.searchsorted(d['time_s'],time);event=manifest['statuses'][key].get('event')
                tables['ThermalSeries'].append(dict(case=case,mode=manifest['series'][key]['mode'],time_s=float(time),
                    Tmin_C=d['Tmin_K'][i]-273.15,Tcenter_C=d['Tcenter_K'][i]-273.15,Tsurface_C=d['radial'][i,0,-1]-273.15,
                    Cmax=d['Cmax'][i],Cmean=d['Cmean'][i],radius_m=d['radius_m'][i],
                    drying_state='DRY_CONTINUED_REAL_STATE' if event and time>=event['report_s'] else 'NOT_YET_DRY',
                    time_source='genuine event-time state' if time in manifest['series'][key].get('exact_event_times_s',[]) else 'actual saved sample'))
    for entry in manifest['counterfactuals']:
        with np.load(root/entry['path']) as data:
            important=[]
            for key in ['control_id','reference_id']:
                event=manifest['statuses'][entry[key]].get('event')
                if event:important.append(event['report_s'])
            indices=np.unique(np.r_[np.arange(0,len(data['time_s']),10),len(data['time_s'])-1,np.flatnonzero(np.isin(data['time_s'],important))])
            for i in indices:tables['Counterfactuals'].append(dict(case=entry['case'],kind=entry['kind'],tail_minutes=entry['tail_minutes'],
                **{k:Export_GetValue(data[k][i]) for k in data.files}))
    return tables


def Export_BuildSupplement(root,manifest,mode):
    from .analysis import Analysis_LoadSeries,Analysis_GetExactState
    from .synthesis import Synthesis_Find
    from .trajectory import Trajectory_GetState,Trajectory_GetNodes,Trajectory_GetInputs
    root=Path(root);wb=Workbook();wb.remove(wb.active)
    config_rows=[]
    for case in ['q1','q23','q4']:
        key=Synthesis_Find(manifest,case,mode)
        if not key:
            config_rows.append(dict(case=case,mode=mode,status='STUDY_REFERENCE_INCOMPLETE',reason='production trajectory unavailable'));continue
        spec=manifest['specs'][key];status=manifest['statuses'][key]
        def Add(quantity,value,unit='',state='RECORDED',source=key):
            config_rows.append(dict(case=case,mode=mode,quantity=quantity,value=value,unit=unit,status=state,source=source))
        Add('drying_time',status.get('drying_time_h'),'h',status['status'])
        units={'R0_m':'m','L_m':'m','T0_K':'K','C0':'kg/kg','h':'W/(m² K)','hm':'m/s','threshold':'kg/kg','t_cap_s':'s'}
        for name,value in spec['physics'].items():Add('configured_global_cap_s' if name=='t_cap_s' else name,value,units[name])
        Add('actual_run_cap_s',spec['schedule'][-1]['t_end'],'s')
        Add('latent_heat_enabled',mode in ['M10','M11']);Add('explicit_moisture_heat_enabled',mode in ['M01','M11'])
        Add('requested_dt_max',spec['dt_max_s'],'s');Add('actual_dt_min',status['minimum_dt'],'s');Add('actual_dt_max',status['maximum_dt'],'s')
        for i,stage in enumerate(spec['schedule'],1):
            Add(f'stage_{i}_Nr',stage['nr'],'cells');Add(f'stage_{i}_start',stage['t_start'],'s');Add(f'stage_{i}_end',stage['t_end'],'s')
        for key_name,field in [('time_verification','time_half_passed'),('spatial_verification','full_schedule_reference_passed')]:
            value=next((c['status'] for c in manifest['checks'] if c['production_id']==key and c[field] is not None),'INCOMPLETE')
            Add(key_name,value,state=value)
        balance=next((a['status'] for a in manifest.get('thermal_audits',[]) if a['experiment_id']==key),'INCOMPLETE')
        Add('effective_balance_verification',balance,state=balance)
        water=spec['water'];Add('water_library_version',water['library_version'],source=water['sources'][1])
        Add('water_path',water['state_path'],source=water['sources'][0])
        Add('water_minimum_temperature',water['temperature_range_K'][0],'K');Add('water_maximum_temperature',water['temperature_range_K'][1],'K')
        Add('hl_interpolation_max_abs_error',water['interpolation_max_abs_J_kg']['hl'],'J/kg')
        Add('Lv_interpolation_max_abs_error',water['interpolation_max_abs_J_kg']['Lv'],'J/kg')
        Add('M00_reference','results/tables/result1.xlsx .. result4.xlsx',source='frozen official workbooks; no new M00 workbook')
    Export_AddTable(wb,'Config_Status',config_rows)
    templates=Baseline_ReadJson(root/'data/input_manifest.json')['templates']
    for q,case,model in [(1,'q1',1),(2,'q23',3),(3,'q23',3),(4,'q4',4)]:
        key=Synthesis_Find(manifest,case,mode)
        if not key:continue
        data=Analysis_LoadSeries(root,manifest['series'][key]);spec=manifest['specs'][key];status=manifest['statuses'][key];event=status.get('event') if q>=3 else None
        end={1:1800.,2:10800.}.get(q,event['report_s'] if event else 259200.)
        times=data['time_s'];keep=(times>0)&(times<=end+1e-8)&(np.abs(times/(1 if q<=2 else 60)-np.round(times/(1 if q<=2 else 60)))<1e-8)
        times=times[keep];values=data['radial'][keep].copy()
        if event and not np.any(np.abs(times-end)<1e-8):
            state,mesh=Analysis_GetExactState(root,spec,end);r,z,nodes=Trajectory_GetNodes(state,end,model,Trajectory_GetInputs(root,spec),mesh,spec)
            positions=np.r_[np.arange(21)*.001,r[-1]];positions[positions>r[-1]+1e-14]=np.nan
            row=np.array([np.interp(positions,r,nodes[f,:,0]) for f in [0,1]])
            times=np.r_[times,end];values=np.concatenate([values,row[None]],axis=0)
        T=values[:,0,:]-273.15;C=values[:,1,:];columns=[round(i*.1,1) for i in range(21)]
        if q==4:columns+=['药材表面']
        else:T=T[:,:21];C=C[:,:21]
        arrays=[T,C] if q<=2 else [C]
        intermediate=root/f'work/studies/experiments/{key}/q{q}_supplement.xlsx'
        Export_WriteWorkbook(root/templates[str(q)]['path'],intermediate,times,arrays,columns)
        saved=load_workbook(intermediate);style_cache={}
        for index,original in enumerate(saved):
            label='Temperature_C' if q<=2 and index==0 else 'Moisture_kg_kg'
            sheet=wb.create_sheet(f'Q{q}_{label}');sheet.freeze_panes='B2';sheet.sheet_view.showGridLines=False
            for row in original:
                for cell in row:
                    new=sheet.cell(cell.row,cell.column,cell.value)
                    if cell.style_id not in style_cache:
                        Export_CopyCellStyle(cell,new);style_cache[cell.style_id]=copy(new._style)
                    else:
                        # These cached IDs have already been registered in wb.
                        new._style=copy(style_cache[cell.style_id])
            for col,dimension in original.column_dimensions.items():
                sheet.column_dimensions[col].width=dimension.width;sheet.column_dimensions[col].hidden=dimension.hidden
            sheet.auto_filter.ref=sheet.dimensions;sheet['A1']='时间 / s\n径向位置 / cm'
            sheet['A1'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);sheet.row_dimensions[1].height=30
        saved.close()
        if q>=3:
            # Original Q3/Q4 templates request moisture only. Their temperatures
            # are separate, explicitly supplementary sheets in this new book.
            reference_sheet=wb[f'Q{q}_Moisture_kg_kg']
            sheet=wb.create_sheet(f'Q{q}_Temperature_C_extension');sheet.append(['时间 / s\n径向位置 / cm']+columns)
            for t,row in zip(times,T):sheet.append([float(t)]+[Export_GetValue(v) for v in row])
            sheet.freeze_panes='B2';sheet.sheet_view.showGridLines=False
            for row in sheet:
                for cell in row:
                    cell._style=copy(reference_sheet.cell(min(cell.row,2),cell.column)._style)
                    if isinstance(cell.value,(int,float)):cell.number_format='0.0000'
            sheet.row_dimensions[1].height=30;sheet.auto_filter.ref=sheet.dimensions
            sheet.column_dimensions['A'].width=29
            for column in range(2,len(columns)+2):sheet.column_dimensions[sheet.cell(1,column).column_letter].width=13
    return Export_SaveVerified(root,wb,root/f'results/studies/thermal/{mode}/supplement.xlsx')


def Studies_Export(root,manifest):
    root=Path(root);outputs=[]
    for name,tables in [('study_summary',manifest['summary_tables']),('study_timeseries',Export_GetTimeseries(root,manifest))]:
        print('STUDY_WORKBOOK '+name,flush=True)
        wb=Workbook();wb.remove(wb.active)
        for title,rows in tables.items():Export_AddTable(wb,title,rows)
        outputs.append(Export_SaveVerified(root,wb,root/f'results/studies/tables/{name}.xlsx'))
    for mode in ['M10','M01','M11']:
        print('STUDY_WORKBOOK '+mode+'/supplement',flush=True)
        outputs.append(Export_BuildSupplement(root,manifest,mode))
    Storage_WriteJson(root/'work/studies/diagnostics/workbook_readback.json',dict(outputs=outputs,
        timeseries_sampling='Actual samples: 60s up to 4h, 600s later, include endpoints/rate extrema/remesh and genuine reintegrated event times. ThermalSeries uses common physical times for every available mode. Full arrays in sealed NPZ; supplements follow official question frequency and own event.',
        official_workbooks='never written'))
    return outputs
