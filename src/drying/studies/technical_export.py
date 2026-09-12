"""Narrow workbook edits and publication receipts for the final technical pass."""
from pathlib import Path
import hashlib
import json
from copy import copy
import numpy as np
from openpyxl import load_workbook
from .baseline import Baseline_ReadJson, Baseline_HashFile
from .export import Export_AddTable, Export_SaveVerified, Export_CopyCellStyle
from .paper_facts import PaperFacts_Build, PaperFacts_WriteMarkdown
from ..storage import Storage_WriteJson


def Technical_GetSheetSignature(sheet):
    digest=hashlib.sha256();styles={}
    for row in sheet:
        for cell in row:
            if cell.style_id not in styles:
                style_text=str((copy(cell.font),copy(cell.fill),copy(cell.border),
                    copy(cell.alignment),copy(cell.protection),cell.number_format))
                styles[cell.style_id]=hashlib.sha256(style_text.encode()).hexdigest()
            digest.update(json.dumps((cell.coordinate,cell.value,styles[cell.style_id]),default=str).encode())
    layout=dict(freeze_panes=sheet.freeze_panes,auto_filter=sheet.auto_filter.ref,
        merged=[str(r) for r in sheet.merged_cells.ranges],
        columns={k:str(v) for k,v in sheet.column_dimensions.items()},
        rows={k:str(v) for k,v in sheet.row_dimensions.items()})
    digest.update(json.dumps(layout,sort_keys=True,default=str).encode())
    return digest.hexdigest()


def Technical_ExportWorkbooks(root, manifest, technical):
    root=Path(root);records=[]
    summary_path=root/'results/studies/tables/study_summary.xlsx'
    book=load_workbook(summary_path)
    protected={s.title:Technical_GetSheetSignature(s) for s in book if s.title!='Geometry_Property_Cross'}
    if 'Geometry_Property_Cross' in book:del book['Geometry_Property_Cross']
    cross=technical['geometry_property_cross'];rows=[dict(r) for r in cross['groups']]
    rows.append(dict(group='I',row_kind='factorial_interaction',**cross['interactions'],
        source_case_id='four source cases in preceding rows',convergence_status='STRUCTURE_DIAGNOSTIC; P4_fixed not independently refined'))
    order=['group','row_kind','drying_completed','drying_time_h','Cmax_24h','Cmax_48h','Cmax_72h',
        'Cmean_24h','Cmean_48h','Cmean_72h','qualified_volume_fraction_24h','qualified_volume_fraction_48h','qualified_volume_fraction_72h',
        'endpoint_time_s','endpoint_Cmax','endpoint_Cmean','source_case_id','convergence_status']
    records_for_sheet=[]
    for row in rows:
        ordered={key:row.get(key) for key in order}
        ordered.update({key:value for key,value in row.items() if key not in ordered})
        labeled={key+(' / kg/kg' if key.startswith(('Cmax_','Cmean_')) else ''):value for key,value in ordered.items()}
        records_for_sheet.append(labeled)
    cross_sheet=Export_AddTable(book,'Geometry_Property_Cross',records_for_sheet)
    endpoint_column=next(cell.column for cell in cross_sheet[1] if str(cell.value).startswith('endpoint_Cmax'))
    for row in range(2,cross_sheet.max_row+1):cross_sheet.cell(row,endpoint_column).number_format='0.000000000'
    for name,signature in protected.items():
        if Technical_GetSheetSignature(book[name])!=signature:raise RuntimeError('UNINTENDED_WORKBOOK_EDIT: '+name)
    records.append(Export_SaveVerified(root,book,summary_path))
    print('TECHNICAL_WORKBOOK study_summary',flush=True)

    series_path=root/'results/studies/tables/study_timeseries.xlsx';book=load_workbook(series_path)
    protected={s.title:Technical_GetSheetSignature(s) for s in book if s.title not in ['Kinetics','Kinetics_Summary']}
    sheet=book['Kinetics'];headers=[cell.value for cell in sheet[1]]
    by_mode={(entry['case'],entry['mode']):(key,entry) for key,entry in technical['kinetics'].items()}
    loaded={}
    columns=[('T_mean_volume','K'),('dTmean_dt','K/s'),('C_mean_volume','kg/kg'),
        ('minus_dCmean_dt','(kg/kg)/s'),('R','m'),('minus_dR_dt','m/s')]
    original_count=next((i for i,key in enumerate(headers) if str(key).startswith('T_mean_volume')),len(headers))
    original_values=[tuple(cell.value for cell in row[:original_count]) for row in sheet]
    for j,(name,unit) in enumerate(columns,start=original_count+1):
        cell=sheet.cell(1,j,name+' / '+unit);Export_CopyCellStyle(sheet.cell(1,1),cell)
        sheet.column_dimensions[cell.column_letter].width=max(19,len(name)+5)
    for i in range(2,sheet.max_row+1):
        case=sheet.cell(i,headers.index('case')+1).value;mode=sheet.cell(i,headers.index('mode')+1).value
        time=float(sheet.cell(i,headers.index('time_s')+1).value);key,entry=by_mode[(case,mode)]
        if key not in loaded:
            with np.load(root/entry['path']) as data:loaded[key]={name:data[name].copy() for name in data.files}
        data=loaded[key];at=int(np.searchsorted(data['time_s'],time))
        # XLSX serializes decimal times; choose the same saved event on either
        # side of its final binary rounding bit, never the following minute.
        candidates=[j for j in [at-1,at] if 0<=j<len(data['time_s'])]
        at=min(candidates,key=lambda j:abs(data['time_s'][j]-time))
        if abs(data['time_s'][at]-time)>1e-8:raise RuntimeError('KINETICS_WORKBOOK_TIME_MISMATCH')
        for j,(name,unit) in enumerate(columns,start=original_count+1):
            value=float(data[name][at]);cell=sheet.cell(i,j,value if np.isfinite(value) else None)
            Export_CopyCellStyle(sheet.cell(i,4),cell);cell.number_format='0.000000' if abs(value)>=1e-4 else '0.0000E+00'
    if original_values!=[tuple(cell.value for cell in row[:original_count]) for row in sheet]:
        raise RuntimeError('UNINTENDED_KINETICS_VALUE_EDIT')
    sheet.auto_filter.ref=sheet.dimensions
    if 'Kinetics_Summary' in book:del book['Kinetics_Summary']
    timing=technical['kinetics_summary'];timing_rows=[]
    quantities=['observation_start_s','observation_end_s','t_T_peak','vT_peak','t_C_peak','vC_peak',
        'T_peak_at_window_boundary','C_peak_at_window_boundary','R_peak_interval_start','R_peak_interval_end','R_peak_intervals_s','vR_peak',
        'delta_R_to_T_interval_start_s','delta_R_to_T_interval_end_s','delta_R_to_C_interval_start_s','delta_R_to_C_interval_end_s','delta_T_to_C',
        't_T50','t_C50','t_R50','delta_R50_to_T50','delta_R50_to_C50','delta_T50_to_C50']
    for name in quantities:
        unit='K/s' if name=='vT_peak' else '(kg/kg)/s' if name=='vC_peak' else 'm/s' if name=='vR_peak' else 'boolean' if name.endswith('boundary') else 's'
        timing_rows.append(dict(quantity=name,value=timing[name],unit=unit,source_case_id=timing['source_case_id'],
            derivative_method=timing['derivative_method'],smoothing_method=timing['smoothing_method']))
    for name,value in timing['half_response_details'].items():
        timing_rows.append(dict(quantity=name+'50_bracket_s',value=value['bracket_s'],unit='s',source_case_id=timing['source_case_id'],
            derivative_method=value['method'],smoothing_method='none'))
    Export_AddTable(book,'Kinetics_Summary',timing_rows)
    for name,signature in protected.items():
        if Technical_GetSheetSignature(book[name])!=signature:raise RuntimeError('UNINTENDED_WORKBOOK_EDIT: '+name)
    records.append(Export_SaveVerified(root,book,series_path))
    print('TECHNICAL_WORKBOOK study_timeseries',flush=True)
    return records


def Technical_CarryRenderReceipt(root, manifest):
    root=Path(root);folder=root/'work/studies/technical'
    original=Baseline_ReadJson(folder/'pre_extension_render.json');protected=Baseline_ReadJson(folder/'protected_existing.json')
    retained=[]
    for item in original['files']:
        if item['path'] not in protected:continue
        if Baseline_HashFile(root/item['path'])!=item['sha256']:raise RuntimeError('EXISTING_FIGURE_MODIFIED')
        retained.append(dict(item,carried_forward_from_manifest_seal=original['manifest_seal'],
            manifest_seal=manifest['seal'],preservation_reason='requested byte-identical preservation; all original scientific payload hashes checked unchanged'))
    report=dict(manifest_seal=manifest['seal'],status='PARTIAL_RENDER',files=retained,failures=[],
        numerical_data='unchanged previous scientific payloads plus independently sealed technical extension',
        carry_forward_scope='seven existing PNG and two existing study GIF; only kinetics/cross figures require rendering')
    Storage_WriteJson(root/'work/studies/diagnostics/render_manifest.json',report)


def Technical_Publish(root, manifest, technical):
    from ..storage import Storage_WriteJson
    root=Path(root)
    # Check factual consistency before updating any of the public output files.
    facts=PaperFacts_Build(root,manifest,technical,publish=False)
    outputs=Technical_ExportWorkbooks(root,manifest,technical)
    by_path={item['path']:item for item in outputs}
    manifest['workbooks']=[by_path.get(item['path'],item) for item in manifest['workbooks']]
    readback_path=root/'work/studies/diagnostics/workbook_readback.json';readback=Baseline_ReadJson(readback_path)
    readback['outputs']=manifest['workbooks'];readback['technical_edits']='only Geometry_Property_Cross and Kinetics columns/Kinetics_Summary; existing sheets/columns verified unchanged'
    Storage_WriteJson(readback_path,readback)
    manifest['summary_tables']['Geometry_Property_Cross']=technical['geometry_property_cross']['groups']
    manifest['summary_tables']['Kinetics_Summary']=[technical['kinetics_summary']]
    path='work/studies/technical/plot_payload/technical_manifest.json'
    manifest['technical_extension']=dict(path=path,sha256=Baseline_HashFile(root/path),status=technical['status'],
        cross_validation_status=technical['cross_validation']['status'],new_1d_trajectories=len(technical['cross_validation']['specs']))
    files={entry['path']:entry for entry in manifest['data_files']}
    for item in technical['payload_files']:files[item['path']]=item
    files[path]=dict(path=path,sha256=Baseline_HashFile(root/path),purpose='sealed final technical metadata')
    manifest['data_files']=list(files.values())
    for item in manifest['figure_plan']:
        if item['name']=='geometry_control':item.update(name='geometry_property_cross',path='results/studies/figures/05_geometry_property_cross.png',source='four property/geometry groups; independent P3-shrink validation')
        if item['name']=='drying_kinetics':item['source']='frozen M00 curves and volume-weighted response timing'
    manifest.pop('seal',None)
    manifest['seal']=hashlib.sha256(json.dumps(manifest,sort_keys=True,allow_nan=False).encode()).hexdigest()
    Storage_WriteJson(root/'work/studies/plot_payload/study_manifest.json',manifest)
    index_path=root/'results/studies/study_index.json';index=Baseline_ReadJson(index_path)
    index['technical_extension']=dict(study_id=technical['study_id'],status=technical['status'],
        cross_production_id=technical['cross_validation']['specs'][0]['experiment_id'],
        completed_1d_trajectories=3,expected_1d_trajectories=3,refinement2d=technical['refinement2d'],
        paper_facts=['results/paper_facts.json','results/paper_facts.md'])
    Storage_WriteJson(index_path,index)
    validation_path=root/'results/studies/validation_summary.json';validation=Baseline_ReadJson(validation_path)
    validation['technical_extension']=dict(status=technical['status'],cross=technical['cross_validation'],refinement2d=technical['refinement2d'])
    from .mass_report import MassReport_ReadSummary
    mass_summary=MassReport_ReadSummary(root)
    if mass_summary:validation['solver_mass_balance']=mass_summary
    Storage_WriteJson(validation_path,validation)
    Storage_WriteJson(root/'results/paper_facts.json',facts)
    PaperFacts_WriteMarkdown(root,facts)
    Technical_CarryRenderReceipt(root,manifest)
    from .synthesis import Synthesis_RefreshOverview
    Synthesis_RefreshOverview(root)
    print('TECHNICAL_PUBLICATION_READY',flush=True)
