"""Freeze explicit production references; never duplicate the large solver cache."""
import hashlib
import json
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
import numpy as np
from PIL import Image
from openpyxl import load_workbook
from ..storage import Storage_WriteJson


def Baseline_ReadJson(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def Baseline_HashFile(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream,'sha256').hexdigest()


def Baseline_Freeze(root):
    from ..cases import Case_ReadStatus, Case_LoadConfig, Case_GetSourceHash
    root=Path(root); folder=root/'work/baseline_snapshot'; path=folder/'baseline_manifest.json'
    if path.exists():
        return Baseline_ReadJson(path)
    folder.mkdir(parents=True,exist_ok=True)
    status=Baseline_ReadJson(root/'results/status.json')
    validation=Baseline_ReadJson(root/'work/validation/summary.json')
    selected={}
    for case in ['q1','q23','q4']:
        for dim in [1,2]:
            case_id=validation[f'{case}_{dim}d']['selected_id']
            selected[f'{case}_{dim}d']=Case_ReadStatus(root,case_id)
    protected=set()
    for source in selected.values():
        ids=[source['case_id']]+[s['case_id'] for s in source.get('stages',[])]
        for case_id in ids:
            protected.update(p for p in (root/'work/cache'/case_id).glob('*') if p.is_file())
    for name in ['data/inputs.npz','data/input_manifest.json','configs/default.toml','configs/stage_schedule.json',
                 'configs/assumptions.json','results/status.json','results/overview.md','work/validation/summary.json',
                 'work/diagnostics/export_manifest.json']:
        protected.add(root/name)
    for pattern in ['*.npz','*.json']:
        protected.update((root/'work/validation/mesh_profiles').glob(pattern))
    for name in ['materials.py','boundaries.py','operators.py','rk4.py','sampling.py','inputs.py','events.py','geometry.py',
                 'cases.py','stages.py']:
        protected.add(root/'src/drying'/name)
    workbooks=[]
    for q in range(1,5):
        p=root/f'results/tables/result{q}.xlsx'; protected.add(p)
        wb=load_workbook(p,read_only=True,data_only=True); sheets=[]
        for sheet in wb:
            digest=hashlib.sha256(); count=0; extrema=[float('inf'),float('-inf')]
            for row in sheet.values:
                digest.update(json.dumps(row,ensure_ascii=False,default=str).encode())
                for value in row:
                    if isinstance(value,(int,float)):
                        count+=1;extrema=[min(extrema[0],value),max(extrema[1],value)]
            sheets.append(dict(name=sheet.title,rows=sheet.max_row,columns=sheet.max_column,
                               numeric_cells=count,numeric_range=extrema,value_hash=digest.hexdigest()))
        wb.close()
        workbooks.append(dict(path=p.relative_to(root).as_posix(),sha256=Baseline_HashFile(p),sheets=sheets))
        backup=folder/'official_files'/p.name;backup.parent.mkdir(exist_ok=True)
        shutil.copy2(p,backup)
    images=[]
    for p in sorted((root/'results').rglob('*')):
        if p.suffix.lower() not in ['.png','.gif'] or 'studies' in p.parts:continue
        with Image.open(p) as im: size=list(im.size);frames=getattr(im,'n_frames',1)
        images.append(dict(path=p.relative_to(root).as_posix(),sha256=Baseline_HashFile(p),size_px=size,frames=frames))
        protected.add(p)
    protected.update((root/'results').glob('q*/q*_summary.json'))
    records={p.relative_to(root).as_posix():dict(sha256=Baseline_HashFile(p),size=p.stat().st_size) for p in sorted(protected)}
    specification=dict(input=Baseline_ReadJson(root/'data/input_manifest.json'),config=Case_LoadConfig(root),
                       numerical_core_hash=Case_GetSourceHash(root),selected=selected,workbooks=workbooks)
    baseline_id=hashlib.sha256(json.dumps(specification,sort_keys=True).encode()).hexdigest()
    try:
        commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=root,text=True,stderr=subprocess.DEVNULL).strip()
        dirty=subprocess.check_output(['git','status','--porcelain'],cwd=root,text=True,stderr=subprocess.DEVNULL).splitlines()
    except (OSError,subprocess.CalledProcessError):
        # Release recomputation has reviewable source, but requires no Git install.
        commit='release-source-copy'; dirty=[]
    result=dict(schema_version=1,baseline_id=baseline_id,created_at=datetime.now().astimezone().isoformat(),
        commit=commit,working_tree_dirty=bool(dirty),working_tree_changes=dirty,**specification,
        status=status,validation=validation,images=images,protected_files=records,
        policy='Frozen M00 official model; extension evidence never replaces official acceptance or data.')
    Storage_WriteJson(path,result)
    print(f'BASELINE_FROZEN {baseline_id[:12]}: {len(records)} protected files',flush=True)
    return result


def Baseline_Check(root):
    root=Path(root); baseline=Baseline_ReadJson(root/'work/baseline_snapshot/baseline_manifest.json')
    migration=root/'work/baseline_snapshot/gif_migration.json'
    paths=Baseline_ReadJson(migration).get('paths',{}) if migration.exists() else {}
    changed=[]
    for name,record in baseline['protected_files'].items():
        p=root/paths.get(name,name)
        if not p.is_file() or Baseline_HashFile(p)!=record['sha256']:
            from ..auxiliary2d import Auxiliary_CheckMetadata
            if not p.is_file() or not Auxiliary_CheckMetadata(root,name,record['sha256']):changed.append(name)
    if changed:raise RuntimeError('BASELINE_OUTPUT_MODIFIED: '+', '.join(changed))
    return dict(status='PASS',baseline_id=baseline['baseline_id'],protected_files=len(baseline['protected_files']),
                official_workbooks_unchanged=True,production_arrays_and_events_unchanged=True)
