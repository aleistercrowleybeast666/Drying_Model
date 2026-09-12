"""Final extension acceptance, including workbook scope and preserved old artifacts."""
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from .baseline import Baseline_ReadJson, Baseline_HashFile
from .technical_contract import Technical_ReadPayload
from ..storage import Storage_WriteJson


def Technical_CheckWorkbookScope(root, filename):
    from .technical_export import Technical_GetSheetSignature
    root=Path(root);original=root/'work/studies/technical/original_workbooks'/filename
    final=root/'results/studies/tables'/filename
    before=load_workbook(original);after=load_workbook(final);checked=[]
    try:
        for old in before:
            current=after[old.title]
            if old.title!='Kinetics':
                if Technical_GetSheetSignature(old)!=Technical_GetSheetSignature(current):
                    raise RuntimeError('ORIGINAL_SHEET_CHANGED: '+filename+'/'+old.title)
            else:
                if current.max_column!=old.max_column+6 or current.max_row!=old.max_row:
                    raise RuntimeError('KINETICS_COLUMN_SCOPE_CHANGED')
                for row in old:
                    for cell in row:
                        peer=current.cell(cell.row,cell.column)
                        if cell.value!=peer.value or cell._style!=peer._style:
                            raise RuntimeError('ORIGINAL_KINETICS_CELL_CHANGED: '+cell.coordinate)
                if old.freeze_panes!=current.freeze_panes:raise RuntimeError('KINETICS_LAYOUT_CHANGED')
                for name,dimension in old.column_dimensions.items():
                    if dimension.width!=current.column_dimensions[name].width:raise RuntimeError('KINETICS_OLD_COLUMN_WIDTH_CHANGED')
            checked.append(old.title)
        expected={'Geometry_Property_Cross'} if filename=='study_summary.xlsx' else {'Kinetics_Summary'}
        if set(after.sheetnames)-set(before.sheetnames)!=expected:raise RuntimeError('UNREQUESTED_WORKSHEET_ADDED')
    finally:before.close();after.close()
    return dict(path=final.relative_to(root).as_posix(),sha256=Baseline_HashFile(final),
        original_sha256=Baseline_HashFile(original),original_sheets_preserved=checked,status='PASS')


def Technical_CheckFinal(root, manifest):
    root=Path(root);technical=Technical_ReadPayload(root,manifest);validation=technical['cross_validation'];problems=[]
    protected=Baseline_ReadJson(root/'work/studies/technical/protected_existing.json')
    changed=[name for name,digest in protected.items() if not (root/name).is_file() or Baseline_HashFile(root/name)!=digest]
    if changed:problems.append('old study data/figures changed: '+', '.join(changed))
    if validation['status']!='PASS':problems.append('P3 shrink own validation '+validation['status'])
    assert validation['inherited_pass'] is False and len(validation['checks'])==2
    assert len(validation['specs'])==len(validation['statuses'])==3
    assert len({s['experiment_id'] for s in validation['specs']})==3
    for spec,status in zip(validation['specs'],validation['statuses']):
        assert spec['fingerprint']==status['fingerprint'] and spec['case']=='q23' and spec['material_appendix']==3 and spec['shrink']
        assert status['complete'] and status['full_history_verified'] and status['simulated_time_s']==259200.
        first=sorted((root/'work/studies/experiments'/spec['experiment_id']).glob('chunk_*.npz'))[0]
        with np.load(first) as data:
            assert data['time_s'][0]==0. and np.all(data['fields'][0,0]==spec['physics']['T0_K']) and np.all(data['fields'][0,1]==spec['physics']['C0'])
    for check in validation['checks']:
        assert check['production_id']==validation['specs'][0]['experiment_id']
        assert check['reference_pairing_verified'] and check['full_history_verified']
        assert all(row['source']=='genuine reintegration to exact common physical time' for row in check['events'])
        if check['status']!='PASS':problems.append('own comparison '+check['reference_id']+' '+check['status'])
    assert all(record['status']=='PASS' for record in validation['transfers'])
    assert all(record['status'] in ['PASS','NOT_DRY_WITHIN_72H'] for record in validation['events'])
    geometry=technical['geometry_property_cross'];rows=geometry['groups']
    assert {r['group'] for r in rows}=={'P3_fixed','P3_shrink','P4_fixed','P4_shrink'}
    if any(row['drying_time_h'] is None for row in rows):assert geometry['interactions']['drying_time_h'] is None
    for row in rows:
        for h in [24,48,72]:assert 0<=row[f'qualified_volume_fraction_{h}h']<=1.
    facts=Baseline_ReadJson(root/'results/paper_facts.json');official=Baseline_ReadJson(root/'results/status.json')['questions']
    for q,row in facts['official'].items():
        current=official[q.lower()]
        assert row['source_case_id']==current['one_id'] and row['drying_time_s']==current['drying_time']
        assert abs(row['key_values']['Cmax']-current['endpoint_values']['max_moisture'])<=1e-12
        assert row['convergence_status']==dict(time=current['time_convergence_passed'],spatial=current['spatial_convergence_passed'])
    assert not facts['source_conflicts'] and (root/'results/paper_facts.md').is_file()
    timing=technical['kinetics_summary']
    assert timing['observation_end_s']==official['q4']['drying_time']
    assert all(entry['maximum_Cmean_consistency_difference']<=1e-12 for entry in technical['kinetics'].values())
    assert all(not row['independent_grid_certification'] for row in technical['refinement2d'].get('cases',{}).values())
    assert all(row['grid']==[60,188] for row in technical['refinement2d'].get('cases',{}).values())
    books=[Technical_CheckWorkbookScope(root,name) for name in ['study_summary.xlsx','study_timeseries.xlsx']]
    result=dict(status='FAIL' if problems else 'PASS',problems=problems,protected_existing_study_files=len(protected),
        old_scientific_arrays_and_artifacts_preserved=not changed,workbooks=books,new_1d_trajectories=3,own_comparisons=2,
        original_innovations_preserved=['B1/B2','B3','B4','B5','B6','B7','B8','B10','B11','B12','M00/M10/M01/M11'],
        B11_event_method='original Event_Locate numerical core remains frozen; own comparison samples use genuine reintegration',
        refinement2d_status=technical['refinement2d']['status'],paper_facts_source_consistency='PASS')
    Storage_WriteJson(root/'work/studies/technical/validation/acceptance.json',result)
    print('TECHNICAL_ACCEPTANCE '+result['status'],flush=True)
    return result
