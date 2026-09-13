"""Minimal A-only process owner. Shares numerical files, never GUI orchestration."""
import argparse
from datetime import datetime
from enum import IntEnum
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import traceback
import uuid
from .storage import Storage_WriteJson


class OfficialRunResult(IntEnum):
    COMPLETE=0
    FAILED=1
    STOPPED=2


def Official_Read(path):
    try:return json.loads(Path(path).read_text(encoding='utf-8'))
    except FileNotFoundError:return {}


def Official_Hash(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def Official_GetCode(root):
    candidate=Path(root)/'dependencies/application/code'
    return candidate if (candidate/'src').is_dir() else Path(root)


def Official_GetData(root):
    candidate=Path(root)/'dependencies/application/data'
    return candidate if candidate.is_dir() else Path(root)/'data'


def Official_GetSeal(value):
    return hashlib.sha256(json.dumps({k:v for k,v in value.items() if k!='seal'},sort_keys=True,allow_nan=False).encode()).hexdigest()


def Official_GetIdentity(root):
    code=Official_GetCode(root);data=Official_GetData(root)
    modules=['materials','boundaries','operators','rk4','sampling','inputs','events','geometry','cases','stages','table_solver']
    files={f'src/drying/{name}.py':Official_Hash(code/f'src/drying/{name}.py') for name in modules}
    for name in ['default.toml','stage_schedule.json','stage_solver_reference.json']:
        files['configs/'+name]=Official_Hash(code/'configs'/name)
    files.update({p.relative_to(code).as_posix():Official_Hash(p) for p in (code/'configs/frozen_mesh').glob('*') if p.is_file()})
    return dict(input_hash=Official_Read(data/'input_manifest.json')['hash'],input_arrays_sha256=Official_Hash(data/'inputs.npz'),files=files)


def Official_CacheCheck(root,case,identity):
    manifest=Official_Read(Path(root)/'results/data/official_manifest.json')
    if manifest.get('seal')!=Official_GetSeal(manifest):return None
    entry=manifest.get('entries',{}).get(case,{})
    if entry.get('identity')!=identity or entry.get('seal')!=Official_GetSeal(entry) or not entry.get('complete'):return None
    for name,digest in entry['files'].items():
        path=(Path(root)/name).resolve()
        if not path.is_relative_to(Path(root).resolve()) or not path.is_file() or Official_Hash(path)!=digest:return None
    return entry


def Official_CopyFile(source,destination):
    destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True)
    if destination.is_file() and Official_Hash(source)==Official_Hash(destination):return str(destination)
    temporary=destination.with_name(destination.name+'.'+uuid.uuid4().hex+'.copy')
    try:shutil.copy2(source,temporary);os.replace(temporary,destination)
    finally:temporary.unlink(missing_ok=True)
    return str(destination)


def Official_Prepare(root,case,force=False):
    workspace=Path(root)/'work/runtime'/case;workspace.mkdir(parents=True,exist_ok=True)
    if force and (workspace/'work').exists():
        source=(workspace/'work').resolve();target=(Path(root)/'work/archived'/f'{case}_{time.time_ns()}').resolve()
        if not source.is_relative_to((Path(root)/'work/runtime').resolve()) or not target.is_relative_to((Path(root)/'work/archived').resolve()):raise ValueError('OFFICIAL_PATH_OUTSIDE_PACKAGE')
        target.parent.mkdir(parents=True,exist_ok=True);shutil.move(str(source),str(target))
    code=Official_GetCode(root)
    for name in ['src','configs']:
        shutil.copytree(code/name,workspace/name,dirs_exist_ok=True,copy_function=Official_CopyFile,
            ignore=shutil.ignore_patterns('__pycache__','*.pyc'))
    shutil.copytree(Official_GetData(root),workspace/'data',dirs_exist_ok=True,copy_function=Official_CopyFile)
    from .frozen_mesh import Frozen_PrepareCase
    Frozen_PrepareCase(workspace,case)
    return workspace


def Official_Readback(root,case):
    import numpy as np
    from openpyxl import load_workbook
    from .export import Export_Readback
    root=Path(root);inputs=Official_Read(root/'data/input_manifest.json')
    for q in dict(q1=[1],q23=[2,3],q4=[4])[case]:
        with np.load(root/f'work/cache/exports/result{q}_full_precision.npz') as data:
            arrays=[data['temperature_C'],data['moisture']] if q<=2 else [data['moisture']]
            columns=[round(x*.1,1) for x in range(21)]+(['药材表面'] if q==4 else [])
            book=load_workbook(root/inputs['templates'][str(q)]['path'],read_only=True);names=book.sheetnames;book.close()
            Export_Readback(root/f'results/tables/result{q}.xlsx',names,data['time_s'],arrays,columns)


def Official_RunCase(root,case):
    from .table_solver import Table_SolveSchedule
    from .table_reference import Table_SealProduction,Table_CheckReference
    from .export import Export_Run
    from .diagnostics import Diagnostics_Open
    root=Path(root);workspace=root/'work/runtime'/case
    Diagnostics_Open(workspace)
    identity=Official_GetIdentity(root);cached=Official_CacheCheck(root,case,identity)
    if cached:
        status=cached['production']
        for q in dict(q1=[1],q23=[2,3],q4=[4])[case]:Official_CopyFile(root/f'results/tables/result{q}.xlsx',workspace/f'results/tables/result{q}.xlsx')
        Official_Readback(workspace,case)
    else:
        status=Table_SolveSchedule(workspace,case)
        Table_SealProduction(workspace,case,status)
        Export_Run(workspace,case,source_case_id=status['case_id'])
        Table_CheckReference(workspace,case,status)
    value=dict(case=case,production=status,cache_reused=bool(cached),status='PASS')
    Storage_WriteJson(root/f'work/receipts/{case}.json',value)
    return OfficialRunResult.COMPLETE


def Official_Publish(root,case,identity):
    root=Path(root);workspace=root/'work/runtime'/case;record=Official_Read(root/f'work/receipts/{case}.json');status=record['production']
    files=[]
    for row in [status,*status['stages']]:
        folder=workspace/'work/cache'/row['case_id'];files += [p for p in folder.iterdir() if p.suffix=='.npz' or p.name=='status.json']
    for q in dict(q1=[1],q23=[2,3],q4=[4])[case]:
        target=root/f'results/tables/result{q}.xlsx';Official_CopyFile(workspace/f'results/tables/result{q}.xlsx',target)
        files += [target,workspace/f'work/cache/exports/result{q}_full_precision.npz']
    entry=dict(identity=identity,complete=True,production=status,files={p.relative_to(root).as_posix():Official_Hash(p) for p in files})
    entry['seal']=Official_GetSeal(entry);path=root/'results/data/official_manifest.json';manifest=Official_Read(path)
    if manifest.get('identity')!=identity:manifest={}
    entries=manifest.get('entries',{});entries[case]=entry
    manifest=dict(schema_version=1,identity=identity,entries=entries,complete=all(c in entries for c in ['q1','q23','q4']),purpose='official_tables_only')
    manifest['seal']=Official_GetSeal(manifest);Storage_WriteJson(path,manifest)
    return record


def Official_GetCommand(root,case):
    root=Path(root)
    command=[sys.executable] if getattr(sys,'frozen',False) else [sys.executable,str(Official_GetCode(root)/'official_recompute.py')]
    return [*command,'--worker',case]


def Official_Run(root,cases,workers='auto',force=False):
    import psutil
    from .judge_schedule import Schedule_GetResources,Schedule_CanStart,Schedule_GetMemoryFailure
    root=Path(root);identity=Official_GetIdentity(root);resources=Schedule_GetResources(workers)
    previous=Official_Read(root/'logs/timings.json');peaks={r['case']:r.get('peak_rss_mb',0) for r in previous.get('cases',[])}
    pending=list(cases);running={};records=[];start=time.perf_counter();last=0.
    env=dict(os.environ,DRYING_OFFICIAL_ROOT=str(root),PYTHONUNBUFFERED='1',PYTHONIOENCODING='utf-8')
    flags=subprocess.CREATE_NO_WINDOW|subprocess.BELOW_NORMAL_PRIORITY_CLASS if os.name=='nt' else 0
    try:
        while pending or running:
            for case in list(pending):
                job=dict(key=case,memory_estimate_mb=max(384.,peaks.get(case,0)*1.3))
                available=psutil.virtual_memory().available/2**20
                if not Schedule_CanStart(job,[r['job'] for r in running.values()],resources['cpu_tokens'],resources['memory_budget_mb'],available,resources['reserve_mb']):continue
                Official_Prepare(root,case,force);(root/f'work/receipts/{case}.json').unlink(missing_ok=True)
                stream=(root/f'logs/{case}.log').open('w',encoding='utf-8');process=subprocess.Popen(Official_GetCommand(root,case),cwd=root,env=env,stdout=stream,stderr=subprocess.STDOUT,creationflags=flags)
                running[case]=dict(process=process,log=stream,job=job,start=time.perf_counter(),peak=0.);pending.remove(case)
            for case,entry in list(running.items()):
                process=entry['process']
                try:entry['peak']=max(entry['peak'],psutil.Process(process.pid).memory_info().rss/2**20)
                except psutil.NoSuchProcess:pass
                if process.poll() is None:continue
                entry['log'].close()
                if process.returncode:raise RuntimeError('OFFICIAL_CASE_FAILED: '+case+'；详情 logs/'+case+'.log')
                row=Official_Publish(root,case,identity);row.update(wall_s=time.perf_counter()-entry['start'],peak_rss_mb=entry['peak'])
                records.append(row);del running[case]
            if pending and not running:
                available=psutil.virtual_memory().available/2**20;job=dict(key=pending[0],memory_estimate_mb=max(384.,peaks.get(pending[0],0)*1.3))
                resources['memory_budget_mb']=max(resources['memory_budget_mb'],available-resources['reserve_mb'])
                if not Schedule_CanStart(job,[],resources['cpu_tokens'],resources['memory_budget_mb'],available,resources['reserve_mb']):raise RuntimeError(Schedule_GetMemoryFailure(job,available,resources))
            now=time.perf_counter()
            if now-last>=2 or not pending and not running:
                event=dict(completed=len(records),total=len(cases),running=list(running),elapsed_s=now-start)
                message=f'已完成 {len(records)}/{len(cases)}；运行中：'+('、'.join(running) or '无')+f'；用时 {(now-start)/60:.1f} min'
                print(message,flush=True)
                with (root/'logs/progress.jsonl').open('a',encoding='utf-8') as stream:stream.write(json.dumps(event,ensure_ascii=False)+'\n')
                with (root/'logs/console.log').open('a',encoding='utf-8') as stream:stream.write(message+'\n')
                last=now
            if pending or running:time.sleep(.2)
        for row in records:
            if row['case'] in ['q23','q4']:print(('Q3' if row['case']=='q23' else 'Q4')+f" = {row['production']['drying_time_h']:.4f} h",flush=True)
        return OfficialRunResult.COMPLETE
    finally:
        for entry in running.values():
            process=entry['process']
            if process.poll() is None:
                try:
                    parent=psutil.Process(process.pid)
                    for child in parent.children(recursive=True):child.kill()
                    parent.kill();process.wait(timeout=5)
                except psutil.NoSuchProcess:pass
            entry['log'].close()
        Storage_WriteJson(root/'logs/timings.json',dict(cases=records,total_wall_s=time.perf_counter()-start,resources=resources))


def Official_Main(root,argv=None):
    parser=argparse.ArgumentParser(description='仅复算四份正式 Excel；不包含 GUI、绘图、创新实验或完整验证。')
    parser.add_argument('--official',nargs='+',choices=['q1','q23','q4']);parser.add_argument('--workers',default='auto',choices=['auto','1','2','3','4','5','6'])
    parser.add_argument('--worker',choices=['q1','q23','q4'],help=argparse.SUPPRESS)
    for flag in ['force-data','dry-run','verify','runtime-check']:parser.add_argument('--'+flag,action='store_true')
    args=parser.parse_args(argv);root=Path(root).resolve()
    if args.worker:
        try:return Official_RunCase(root,args.worker)
        except Exception:traceback.print_exc();return OfficialRunResult.FAILED
    cases=args.official or ['q23','q4','q1']
    if args.dry_run:print(json.dumps(dict(cases=cases,scope='A only',pde_solves=0),ensure_ascii=False));return OfficialRunResult.COMPLETE
    if args.verify:
        manifest=Official_Read(root/'package_manifest.json')
        if not manifest:raise RuntimeError('PACKAGE_MANIFEST_MISSING')
        for name,digest in manifest['files'].items():
            target=(root/name).resolve()
            if not target.is_relative_to(root) or Official_Hash(target)!=digest:raise RuntimeError('PACKAGE_HASH_MISMATCH: '+name)
        print('PACKAGE_VERIFY: PASS (no numerical solve)');return OfficialRunResult.COMPLETE
    if args.runtime_check:
        import numpy,numba,llvmlite,openpyxl,psutil
        print(json.dumps({m.__name__:m.__version__ for m in [numpy,numba,llvmlite,openpyxl,psutil]}));return OfficialRunResult.COMPLETE
    import psutil
    (root/'logs').mkdir(parents=True,exist_ok=True);(root/'work').mkdir(parents=True,exist_ok=True);lock=root/'work/official.lock'
    if lock.exists():
        previous=Official_Read(lock)
        try:
            if abs(psutil.Process(previous['pid']).create_time()-previous['created_at'])<.01:raise RuntimeError('本目录已有四表计算正在运行。')
        except psutil.NoSuchProcess:pass
        lock.unlink()
    descriptor=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
    os.write(descriptor,json.dumps(dict(pid=os.getpid(),created_at=psutil.Process().create_time())).encode());os.close(descriptor)
    try:return Official_Run(root,cases,args.workers,args.force_data)
    except KeyboardInterrupt:return OfficialRunResult.STOPPED
    except Exception:
        detail=traceback.format_exc();print(detail,flush=True)
        with (root/'logs/console.log').open('a',encoding='utf-8') as stream:stream.write(detail)
        return OfficialRunResult.FAILED
    finally:lock.unlink(missing_ok=True)
